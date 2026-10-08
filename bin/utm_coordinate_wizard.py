"""
UTM Coordinate Wizard (lecturer / student builds)
Malaysia Coordinate Conversion Tool

EA original edition: bin/ea_coordinate_wizard.py
Regenerate EA after logic changes: python bin/tools/_gen_ea_wizard.py
"""

import win_boot  # noqa: F401 — before tkinter (Windows taskbar icon)

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import math, csv, os, sys, zipfile

from app_paths import edition, is_lecturer, is_student
from ui_theme import (
    THEMES, load_theme_mode, save_theme_mode,
    header_logo_image, apply_theme_widgets, _collect_skip,
)

# Student / public builds hide technical detail; lecturer & dev show full UI.
# End-user public portable only (EA_PUBLIC_BUILD=1). Classroom student exe keeps full UI.
PUBLIC_BUILD = os.environ.get("EA_PUBLIC_BUILD") == "1"
LECTURER_MODE = is_lecturer()

# ─────────────────────────────────────────────────────────────────────────────
#  ELLIPSOIDS
# ─────────────────────────────────────────────────────────────────────────────
MOD_EVEREST = {"a": 6377304.063,  "f": 1 / 300.8017}
GRS80       = {"a": 6378137.0,    "f": 1 / 298.257222101}
WGS84       = {"a": 6378137.0,    "f": 1 / 298.257223563}

DATUM_OPTIONS = {
    "PMGSN94 (WGS84)":          ("PMGSN94", WGS84),
    "GDM2009 (GRS80)":          ("GDM2009", GRS80),
    "GDM2000 (GRS80)":          ("GDM2000", GRS80),
    "MRT48 (Modified Everest)": ("MRT48", MOD_EVEREST),
}
DATUM_LABELS = list(DATUM_OPTIONS.keys())

_CASSINI_DATUM_LABEL = {
    "GDM2000": "GDM2000 (GRS80)",
    "GDM2009": "GDM2009 (GRS80)",
    "MRT48": "MRT48 (Modified Everest)",
}

MRT48_POLY_SCALE = 10000.0
CHAIN_M = 20.11678249        # 1 chain = 20.11678249 m (Peninsular Malaysia)
LINK_M = CHAIN_M / 100.0     # 1 link = 0.2011678249 m (100 links per chain)

# Filled from Malaysia.xml (+ companion files) after RSO math is defined below.
DATUM_TRANSFORMS = {}
ZONES = []
RSO_ZONES = []
ZONE_NAMES = []


def _m_to_chain(m):
    """Grid metres → chains for MRT48 polynomial (P[] origins are in chains)."""
    return m / CHAIN_M


def _chain_to_m(ch):
    """Chains → grid metres."""
    return ch * CHAIN_M


def mrt48_grid_input_to_m(n, e, unit="m"):
    """MRT48 Cassini grid input → metres. unit: 'm' (metric plans) or 'links' (50 links/inch plans)."""
    n, e = float(n), float(e)
    if unit == "links":
        return n * LINK_M, e * LINK_M
    return n, e


def mrt48_grid_m_to_unit(n, e, unit="m"):
    """MRT48 Cassini grid metres → display unit (m or links)."""
    n, e = float(n), float(e)
    if unit == "links":
        return n / LINK_M, e / LINK_M
    return n, e


_MRT48_UNIT_OPTS = (("Metric", "m"), ("Links", "links"))


def mrt48_unit_format_hint(unit="m"):
    if unit == "links":
        return "Scale Format Example = 50 links to an inch"
    return "Scale Format Example = 1:1000"


def _dms(d, m, s=0.0):
    return d + m / 60.0 + s / 3600.0


from datum_transform import (
    datum_geo_transform, transform_available as _helmert_key, set_transforms,
    geo_to_ecef_xyz, canon_datum,
)


def _mrt48_zone_params(zone):
    """Extract MRT48 C_S polynomial / origin parameters from zone P[]."""
    P = zone["P"]
    return dict(StE0=P[0], StN0=P[1], RsoE0=P[2], RsoN0=P[3],
                R1=P[4], R2=P[5],
                A=[P[6], P[8], P[10], P[12], P[14]],
                B=[P[7], P[9], P[11], P[13], P[15]])


def _mrt48_poly_corr(p, X, Y):
    """Biquadratic correction (PDF §4.4) — X,Y already scaled by 1/10000."""
    return (p["R1"] + X * p["A"][0] + Y * p["A"][1] + X * Y * p["A"][2]
            + X * X * p["A"][3] + Y * Y * p["A"][4],
            p["R2"] + X * p["B"][0] + Y * p["B"][1] + X * Y * p["B"][2]
            + X * X * p["B"][3] + Y * Y * p["B"][4])


def cassini_to_mrt48_rso(zone, N_grid, E_grid):
    """MRT48 state Cassini → MRT48 RSO via JUPEM polynomial (PDF §4.4 / p.25).

    Grid I/O is in metres; P[] origins and the polynomial operate in chains.
    """
    if zone["type"] != "C_S":
        raise ValueError("Polynomial Cassini↔RSO applies to MRT48 C_S zones only")
    p = _mrt48_zone_params(zone)
    dN = _m_to_chain(N_grid) - p["StN0"]
    dE = _m_to_chain(E_grid) - p["StE0"]
    X, Y = dN / MRT48_POLY_SCALE, dE / MRT48_POLY_SCALE
    cN, cE = _mrt48_poly_corr(p, X, Y)
    return _chain_to_m(dN + p["RsoN0"] + cN), _chain_to_m(dE + p["RsoE0"] + cE)


def mrt48_rso_to_cassini(zone, N_rso, E_rso):
    """MRT48 RSO → MRT48 state Cassini — two-pass inverse (PDF §4.4 / pp.22–24).

    Grid I/O is in metres; P[] origins and the polynomial operate in chains.
    """
    if zone["type"] != "C_S":
        raise ValueError("Polynomial Cassini↔RSO applies to MRT48 C_S zones only")
    p = _mrt48_zone_params(zone)
    dN = _m_to_chain(N_rso) - p["RsoN0"]
    dE = _m_to_chain(E_rso) - p["RsoE0"]
    X1, Y1 = dN / MRT48_POLY_SCALE, dE / MRT48_POLY_SCALE
    cN, cE = _mrt48_poly_corr(p, X1, Y1)
    N1, E1 = dN - cN, dE - cE
    X2, Y2 = N1 / MRT48_POLY_SCALE, E1 / MRT48_POLY_SCALE
    cN2, cE2 = _mrt48_poly_corr(p, X2, Y2)
    return _chain_to_m(dN - cN2 + p["StN0"]), _chain_to_m(dE - cE2 + p["StE0"])


def _mrt48_rso_zone():
    return next(z for z in RSO_ZONES if z["datum"] == "MRT48")


def mrt48_cassini_to_geo(zone, N_grid, E_grid):
    """MRT48 state Cassini → geographic via RSO (JUPEM: Cassini→RSO, then RSO→geo)."""
    rN, rE = cassini_to_mrt48_rso(zone, N_grid, E_grid)
    return rso_to_geo(_mrt48_rso_zone(), rN, rE)


def geo_to_mrt48_cassini(zone, lat_deg, lon_deg):
    """Geographic → MRT48 state Cassini via RSO (JUPEM: geo→RSO, then RSO→Cassini)."""
    rN, rE = geo_to_rso(_mrt48_rso_zone(), lat_deg, lon_deg)
    return mrt48_rso_to_cassini(zone, rN, rE)


def _paired_gdm_zone(mrt48_zone, gdm_datum="GDM2000"):
    """Return the GDM Cassini zone matching an MRT48 state zone."""
    label = mrt48_zone["name"].split("  (")[0].strip()
    if label.startswith("Pahang"):
        label = "Pahang"
    elif label.startswith("Perak"):
        label = "Perak"
    elif label == "Pinang":
        label = "P. Pinang & S. Perai"
    elif label.startswith("Negeri"):
        label = "N. Sembilan & Melaka"
    for z in ZONES:
        if z["type"] == "C_S2000" and z["datum"] == gdm_datum:
            zl = z["name"].split("  (")[0].strip()
            if label in zl or zl in label:
                return z
    return None


def _datum_name(datum_label):
    return DATUM_OPTIONS[datum_label][0]


def _datum_ellipsoid(datum_label):
    return DATUM_OPTIONS[datum_label][1]


def _zone_datum_blurb(z):
    """One-line datum / ellipsoid description for a Cassini zone (Geoid-style)."""
    ell = "Modified Everest" if z["type"] == "C_S" else "GRS80"
    if PUBLIC_BUILD:
        return f"Datum: {z['datum']}   Ellipsoid: {ell}"
    poly = ("  +  polynomial correction" if z["type"] == "C_S"
            else "  pure Cassini-Soldner")
    return f"Datum: {z['datum']}   Ellipsoid: {ell}{poly}"


def _rso_datum_blurb(z):
    ell = "Modified Everest" if z["datum"] == "MRT48" else "GRS80"
    base = f"Datum: {z['datum']}   Ellipsoid: {ell}"
    if PUBLIC_BUILD:
        return f"{base}   RSO"
    return f"{base}   Hotine Oblique Mercator (RSO)"


def packed_dms_to_deg(val):
    """Decode MAGNET Tools packed DDMMSS.sssss → decimal degrees."""
    sign = -1 if val < 0 else 1
    val  = abs(val)
    s    = val % 100
    val  = int(val) // 100
    m    = val % 100
    d    = val // 100
    return sign * (d + m/60 + s/3600)


# ─────────────────────────────────────────────────────────────────────────────
#  WGEOID04  (Malaysia gravimetric geoid 2004)
#  1-arcminute grid, Peninsular Malaysia (WGeoid04.gsf header: 0–8°N, 98–107°E)
# ─────────────────────────────────────────────────────────────────────────────
import os as _os

_GEOID_DATA = None
_GEOID_LOADED_PATH = None
_GEOID_ROWS = 481
_GEOID_COLS = 541
_GEOID_LAT0 = 0.0
_GEOID_LON0 = 98.0
_GEOID_LAT1 = 8.0
_GEOID_LON1 = 107.0
_GEOID_NODATA = 9999.0


def _resolve_geoid_path(path):
    """Prefer companion .gsf when a legacy .gff path is passed."""
    if not path:
        return path
    if _os.path.isfile(path):
        if path.lower().endswith(".gff"):
            from app_paths import GEOID_GSF_NAME, GEOID_GSF_SEALED, geoid_gsf_path

            folder = _os.path.dirname(path)
            for name in (GEOID_GSF_NAME, GEOID_GSF_SEALED, "WGeoid04.gsf", "wgeoid04.gsf"):
                companion = _os.path.join(folder, name)
                if _os.path.isfile(companion):
                    return companion
            gsf = geoid_gsf_path()
            if _os.path.isfile(gsf):
                return gsf
        return path
    if path.lower().endswith(".gff"):
        from app_paths import geoid_file_path
        return geoid_file_path()
    return path


def _load_geoid_gsf(path):
    global _GEOID_DATA, _GEOID_ROWS, _GEOID_COLS
    global _GEOID_LAT0, _GEOID_LON0, _GEOID_LAT1, _GEOID_LON1
    from geoid_vault import read_geoid_bytes

    text = read_geoid_bytes(path).decode("utf-8", errors="replace")
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    if len(lines) < 7:
        return False
    lat0 = float(lines[0])
    lon0 = float(lines[1])
    lat1 = float(lines[2])
    lon1 = float(lines[3])
    ncol_i = int(float(lines[4]))
    nrow_i = int(float(lines[5]))
    cols = ncol_i + 1
    rows = nrow_i + 1
    data = []
    for ln in lines[6:]:
        if ln == "WGeoid04":
            break
        if ln == "N":
            data.append(_GEOID_NODATA)
        else:
            data.append(float(ln))
    total = rows * cols
    if len(data) < total:
        data += [_GEOID_NODATA] * (total - len(data))
    _GEOID_DATA = data[:total]
    _GEOID_ROWS, _GEOID_COLS = rows, cols
    _GEOID_LAT0, _GEOID_LON0 = lat0, lon0
    _GEOID_LAT1, _GEOID_LON1 = lat1, lon1
    return True


def _load_geoid(path):
    global _GEOID_DATA, _GEOID_LOADED_PATH
    path = _resolve_geoid_path(path)
    if not path or not _os.path.isfile(path):
        return False
    norm = _os.path.normcase(_os.path.abspath(path))
    if _GEOID_DATA is not None and _GEOID_LOADED_PATH == norm:
        return True
    _GEOID_DATA = None
    _GEOID_LOADED_PATH = norm
    low = path.lower()
    from geoid_vault import GEOID_GSF_SEALED

    if low.endswith(".gsf") or _os.path.basename(path).lower() == GEOID_GSF_SEALED.lower():
        return _load_geoid_gsf(path)
    return False


def geoid_undulation(lat, lon, geoid_path):
    """Bilinear interpolation of WGeoid04.
    Returns (N_metres, None) on success, or (None, error_message) on failure."""
    resolved = _resolve_geoid_path(geoid_path)
    if not _load_geoid(resolved):
        if geoid_path and geoid_path.lower().endswith(".gff"):
            return None, (
                f"Use {resolved or 'WGeoid04.gsf'} for geoid height — "
                f"the binary GFF grid cannot be read reliably.")
        return None, f"Geoid file not found: {geoid_path}"
    if not (_GEOID_LAT0 <= lat <= _GEOID_LAT1 and
            _GEOID_LON0 <= lon <= _GEOID_LON1):
        return None, (
            f"Out of WGeoid04 coverage "
            f"(Lat {_GEOID_LAT0}–{_GEOID_LAT1}°N, Lon {_GEOID_LON0}–{_GEOID_LON1}°E). "
            f"Got {lat:.6f}N {lon:.6f}E")
    dlat = (_GEOID_LAT1 - _GEOID_LAT0) / (_GEOID_ROWS - 1)
    dlon = (_GEOID_LON1 - _GEOID_LON0) / (_GEOID_COLS - 1)
    rf = (lat - _GEOID_LAT0) / dlat
    cf = (lon - _GEOID_LON0) / dlon
    r0, c0 = int(rf), int(cf)
    r1 = min(r0 + 1, _GEOID_ROWS - 1)
    c1 = min(c0 + 1, _GEOID_COLS - 1)
    dr, dc = rf - r0, cf - c0

    def _v(r, c):
        val = _GEOID_DATA[r * _GEOID_COLS + c]
        return None if abs(val - _GEOID_NODATA) < 1.0 else val

    v00, v01, v10, v11 = _v(r0, c0), _v(r0, c1), _v(r1, c0), _v(r1, c1)
    if any(x is None for x in (v00, v01, v10, v11)):
        return None, "No geoid data at this location (ocean or edge of grid)"
    N = ((1 - dr) * (1 - dc) * v00 + (1 - dr) * dc * v01 +
         dr * (1 - dc) * v10 + dr * dc * v11)
    return N, None



# ─────────────────────────────────────────────────────────────────────────────
#  GEODETIC CORE
# ─────────────────────────────────────────────────────────────────────────────
def _ellipsoid(zone):
    ell = MOD_EVEREST if zone["type"] == "C_S" else GRS80
    a = ell["a"]; f = ell["f"]
    return a, 1-(1-f)**2


def _M(a, e2, phi):
    c0=1-e2/4-3*e2**2/64-5*e2**3/256
    c1=3*e2/8+3*e2**2/32+45*e2**3/1024
    c2=15*e2**2/256+45*e2**3/1024
    c3=35*e2**3/3072
    return a*(c0*phi-c1*math.sin(2*phi)+c2*math.sin(4*phi)-c3*math.sin(6*phi))


def _poly(P, X, Y):
    """Legacy helper — X,Y are chain deltas; use _mrt48_poly_corr in production."""
    x, y = X / MRT48_POLY_SCALE, Y / MRT48_POLY_SCALE
    dN = P[4] + P[6]*x + P[8]*y + P[10]*x*y + P[12]*x**2 + P[14]*y**2
    dE = P[5] + P[7]*x + P[9]*y + P[11]*x*y + P[13]*x**2 + P[15]*y**2
    return dN, dE


def _origin(zone):
    P = zone["P"]
    if zone["type"] == "C_S":
        # P[0]=StE0 (false easting), P[1]=StN0 (false northing)
        # N0/E0 are the Cassini grid false northing/easting, not the RSO params.
        return (math.radians(zone["lat0"]), math.radians(zone["lon0"]),
                P[1], P[0], 0.0, 0.0)
    return (math.radians(packed_dms_to_deg(P[0])),
            math.radians(packed_dms_to_deg(P[1])),
            P[2], P[3], 0.0, 0.0)


def _cs_fwd(a, e2, phi0, lam0, M0, phi, lam):
    M  = _M(a, e2, phi)
    nu = a/math.sqrt(1-e2*math.sin(phi)**2)
    T  = math.tan(phi)**2
    eta2 = (e2/(1-e2))*math.cos(phi)**2
    A_ = (lam-lam0)*math.cos(phi)
    X = M-M0+nu*math.tan(phi)*(A_**2/2+(5-T+6*eta2)*A_**4/24)
    Y = nu*(A_-T*A_**3/6-(8-T+8*eta2)*T*A_**5/120)
    return X, Y


def _cs_inv(a, e2, phi0, lam0, M0, X, Y):
    """Iterative Newton-Raphson inverse — converges to sub-micrometre."""
    c0=1-e2/4-3*e2**2/64-5*e2**3/256
    mu=(M0+X)/(a*c0)
    e1=(1-math.sqrt(1-e2))/(1+math.sqrt(1-e2))
    # Series solution as initial guess
    phi1=(mu+(3*e1/2-27*e1**3/32)*math.sin(2*mu)
          +(21*e1**2/16-55*e1**4/32)*math.sin(4*mu)
          +(151*e1**3/96)*math.sin(6*mu)
          +(1097*e1**4/512)*math.sin(8*mu))
    nu1=a/math.sqrt(1-e2*math.sin(phi1)**2)
    T1=math.tan(phi1)**2; eta1sq=(e2/(1-e2))*math.cos(phi1)**2
    D=Y/nu1
    phi=phi1-(nu1*math.tan(phi1)/(a*(1-e2)/(1-e2*math.sin(phi1)**2)**1.5))*(D**2/2-(1+3*T1)*D**4/24)
    lam=lam0+(D-(1+2*T1+eta1sq)*D**3/6+(5+6*eta1sq+28*T1-3*eta1sq**2+8*eta1sq+24*T1**2)*D**5/120)/math.cos(phi1)
    # Newton-Raphson refinement
    h = 1e-9
    for _ in range(10):
        Xc, Yc = _cs_fwd(a, e2, phi0, lam0, M0, phi, lam)
        dX = Xc - X;  dY = Yc - Y
        if abs(dX) < 1e-9 and abs(dY) < 1e-9:
            break
        Xph, Yph = _cs_fwd(a, e2, phi0, lam0, M0, phi+h, lam)
        Xlh, Ylh = _cs_fwd(a, e2, phi0, lam0, M0, phi, lam+h)
        J00=(Xph-Xc)/h; J01=(Xlh-Xc)/h
        J10=(Yph-Yc)/h; J11=(Ylh-Yc)/h
        det=J00*J11-J01*J10
        phi -= (J11*dX - J01*dY)/det
        lam -= (-J10*dX + J00*dY)/det
    return phi, lam


def geo_to_cassini(zone, lat_deg, lon_deg):
    a,e2 = _ellipsoid(zone)
    phi0,lam0,N0,E0,dN_off,dE_off = _origin(zone)
    M0  = _M(a,e2,phi0)
    phi = math.radians(lat_deg)
    lam = math.radians(lon_deg)
    X,Y = _cs_fwd(a,e2,phi0,lam0,M0,phi,lam)
    return N0+X, E0+Y


def cassini_to_geo(zone, N_grid, E_grid):
    a,e2 = _ellipsoid(zone)
    phi0,lam0,N0,E0,dN_off,dE_off = _origin(zone)
    M0 = _M(a,e2,phi0)
    # Cassini-Soldner: Northing=Y (meridional), Easting=X (transverse)
    # In the math: _cs_fwd returns (X_math, Y_math) where X_math=meridional, Y_math=transverse
    # So: N_grid (Northing) -> X_math,  E_grid (Easting) -> Y_math
    X_rel = N_grid - N0 - dN_off   # Northing -> meridional
    Y_rel = E_grid - E0 - dE_off   # Easting  -> transverse

    # Polynomial correction form not yet validated — omit for accuracy.
    phi, lam = _cs_inv(a,e2,phi0,lam0,M0, X_rel, Y_rel)

    return math.degrees(phi), math.degrees(lam)


# ─────────────────────────────────────────────────────────────────────────────
#  DMS HELPERS
# ─────────────────────────────────────────────────────────────────────────────
# DMS seconds decimals — ~0.1 mm ground precision at the equator (with 9 DD places)
_DMS_SEC_DECIMALS = 6


def dd_to_dms(dd):
    dd_abs = abs(dd)
    d = int(dd_abs)
    m = int((dd_abs - d) * 60)
    s = ((dd_abs - d) * 60 - m) * 60
    return d, m, s


def format_dms_full(dd, is_lat):
    """Compact DMS value: ° ′ ″ with N/S/E/W; seconds to sub-mm precision."""
    d, m, s = dd_to_dms(dd)
    hemi = (("N", "S") if is_lat else ("E", "W"))[0 if dd >= 0 else 1]
    sec = f"{s:09.{_DMS_SEC_DECIMALS}f}"
    return f"{d}°  {m:02d}′  {sec}″  {hemi}"


def dms_to_dd(d, m, s, hemi):
    dd = float(d) + float(m)/60 + float(s)/3600
    return -dd if hemi in ("S", "W") else dd


def split_coord_pair(text):
    """Split 'a, b' / 'a;b' / tab / space-separated pair into two tokens."""
    if not text:
        return None
    s = text.strip()
    for sep in (",", ";", "\t"):
        if sep in s:
            parts = [p.strip() for p in s.split(sep) if p.strip()]
            if len(parts) >= 2:
                return parts[0], parts[1]
    parts = s.split()
    if len(parts) >= 2:
        return parts[0], parts[1]
    return None


def assign_lat_lon(pair, from_above=None):
    """Pick lat/lon from a pair; |lat|≤90. Uses field hint when both values look like lat."""
    a_s, b_s = pair[0].strip(), pair[1].strip()
    try:
        a, b = float(a_s), float(b_s)
    except ValueError:
        if from_above is not None:
            return (a_s, b_s) if from_above else (b_s, a_s)
        return a_s, b_s

    a_lat = abs(a) <= 90
    b_lat = abs(b) <= 90

    if a_lat and not b_lat:
        return a_s, b_s
    if b_lat and not a_lat:
        return b_s, a_s

    if from_above is not None:
        return (a_s, b_s) if from_above else (b_s, a_s)
    return a_s, b_s


def validate_lat_lon_paste(pair):
    """Return an error message if pair cannot be lat/lon degrees, else None."""
    for s in pair[:2]:
        t = str(s).strip().replace(",", "")
        if not t:
            continue
        try:
            v = float(t)
        except ValueError:
            continue
        if abs(v) > 180:
            return (
                "These values are not latitude/longitude.\n\n"
                f"Value {v:g} exceeds 180°, which is the maximum for longitude.\n"
                "Valid ranges: latitude ±90°, longitude ±180°.\n\n"
                "Did you paste Northing/Easting or other grid coordinates instead?"
            )
    return None


_BATCH_PASTE_HEADER = {
    "name": ("name", "point", "id", "label", "station", "pt"),
    "n_in": ("northing", "north", "n_in", "n", "y"),
    "e_in": ("easting", "east", "e_in", "e", "x"),
    "lat_in": ("lat", "latitude", "lat_in"),
    "lon_in": ("lon", "longitude", "long", "lon_in"),
    "zone_in": ("zone", "utm_zone", "zone_in"),
    "hemi_in": ("hemi", "hemisphere", "hemi_in"),
}


def _batch_norm_header(cell):
    t = cell.strip().lower().split("(")[0].strip().replace(" ", "_")
    return t


def _batch_is_header_row(cells):
    for c in cells:
        t = _batch_norm_header(c)
        if not t:
            continue
        try:
            float(t.replace(",", ""))
            continue
        except ValueError:
            pass
        for aliases in _BATCH_PASTE_HEADER.values():
            if t in aliases:
                return True
    return False


def _batch_header_col_map(headers, keys):
    norm = [_batch_norm_header(h) for h in headers]
    mapping = {}
    for key in keys:
        aliases = _BATCH_PASTE_HEADER.get(key, (key,))
        for i, h in enumerate(norm):
            if h and h in aliases:
                mapping[key] = i
                break
    return mapping


def _batch_header_usable(mapping, field_keys):
    return mapping and all(k in mapping for k in field_keys)

# ─────────────────────────────────────────────────────────────────────────────
def utm_zone_from_lon(lon_deg):
    return int((lon_deg + 180) / 6) + 1

def _utm_lam0(zone_num):
    return math.radians((zone_num - 1) * 6 - 180 + 3)

def geo_to_utm(ell, lat_deg, lon_deg, zone_num=None):
    """Geographic → UTM.  Returns (N, E, zone_num, hemisphere_str)."""
    a = ell["a"]; f = ell["f"]; e2 = 2*f - f**2
    if zone_num is None:
        zone_num = utm_zone_from_lon(lon_deg)
    phi  = math.radians(lat_deg)
    lam  = math.radians(lon_deg)
    lam0 = _utm_lam0(zone_num)
    k0   = 0.9996; FE = 500_000.0
    FN   = 0.0 if lat_deg >= 0 else 10_000_000.0
    Nv   = a / math.sqrt(1 - e2*math.sin(phi)**2)
    T    = math.tan(phi)**2
    C    = (e2/(1-e2))*math.cos(phi)**2
    A_   = (lam - lam0)*math.cos(phi)
    ep2  = e2/(1-e2)
    M    = _M(a, e2, phi)
    M0   = _M(a, e2, 0.0)
    E = FE + k0*Nv*(A_ + (1-T+C)*A_**3/6
        + (5-18*T+T**2+72*C-58*ep2)*A_**5/120)
    N = FN + k0*(M - M0 + Nv*math.tan(phi)*(
        A_**2/2 + (5-T+9*C+4*C**2)*A_**4/24
        + (61-58*T+T**2+600*C-330*ep2)*A_**6/720))
    return N, E, zone_num, ("N" if lat_deg >= 0 else "S")

def utm_to_geo(ell, N_grid, E_grid, zone_num, hemisphere):
    """UTM → Geographic.  Returns (lat_deg, lon_deg)."""
    a = ell["a"]; f = ell["f"]; e2 = 2*f - f**2
    k0 = 0.9996; FE = 500_000.0
    FN = 0.0 if hemisphere == "N" else 10_000_000.0
    lam0 = _utm_lam0(zone_num)
    ep2  = e2/(1-e2)
    M0   = _M(a, e2, 0.0)
    M    = M0 + (N_grid - FN)/k0
    c0   = 1 - e2/4 - 3*e2**2/64 - 5*e2**3/256
    mu   = M/(a*c0)
    e1   = (1 - math.sqrt(1-e2))/(1 + math.sqrt(1-e2))
    phi1 = (mu + (3*e1/2 - 27*e1**3/32)*math.sin(2*mu)
            + (21*e1**2/16 - 55*e1**4/32)*math.sin(4*mu)
            + (151*e1**3/96)*math.sin(6*mu)
            + (1097*e1**4/512)*math.sin(8*mu))
    N1   = a/math.sqrt(1 - e2*math.sin(phi1)**2)
    T1   = math.tan(phi1)**2
    C1   = ep2*math.cos(phi1)**2
    R1   = a*(1-e2)/(1 - e2*math.sin(phi1)**2)**1.5
    D    = (E_grid - FE)/(N1*k0)
    phi  = phi1 - (N1*math.tan(phi1)/R1)*(
        D**2/2 - (5+3*T1+10*C1-4*C1**2-9*ep2)*D**4/24
        + (61+90*T1+298*C1+45*T1**2-252*ep2-3*C1**2)*D**6/720)
    lam  = lam0 + (D - (1+2*T1+C1)*D**3/6
        + (5-2*C1+28*T1-3*C1**2+8*ep2+24*T1**2)*D**5/120)/math.cos(phi1)
    # Newton-Raphson refinement to eliminate series truncation residuals
    for _ in range(4):
        Nc, Ec, _, _ = geo_to_utm(ell, math.degrees(phi), math.degrees(lam), zone_num)
        FN_ = 0.0 if hemisphere == "N" else 10_000_000.0
        dN_ = Nc - N_grid; dE_ = Ec - E_grid
        if abs(dN_) < 1e-10 and abs(dE_) < 1e-10:
            break
        h = 1e-8
        Nph, Eph, _, _ = geo_to_utm(ell, math.degrees(phi+h), math.degrees(lam), zone_num)
        Nlh, Elh, _, _ = geo_to_utm(ell, math.degrees(phi),   math.degrees(lam+h), zone_num)
        J00 = (Nph-Nc)/h; J01 = (Nlh-Nc)/h
        J10 = (Eph-Ec)/h; J11 = (Elh-Ec)/h
        det = J00*J11 - J01*J10
        phi -= (J11*dN_ - J01*dE_)/det
        lam -= (-J10*dN_ + J00*dE_)/det
    return math.degrees(phi), math.degrees(lam)


# ─────────────────────────────────────────────────────────────────────────────
#  RSO — Hotine Oblique Mercator (variant A)
#  Used for Malaysia RSO (GDM2000 / GDM2009), EPSG:3375 / EPSG:3376
# ─────────────────────────────────────────────────────────────────────────────
def _rso_constants(ell, lat_c_deg, lon_c_deg, azimuth_deg, gamma_c_deg, k0):
    """Pre-compute RSO projection constants from definition parameters.

    gamma_c_deg: Angle from Rectified to Skew Grid (EPSG parameter).
                 NOT the same as azimuth_deg; must be supplied explicitly.
                 For all Malaysia RSO variants: 323.130102361111°  (EPSG).
    """
    a = ell["a"]; f = ell["f"]
    e2 = 2*f - f**2
    e  = math.sqrt(e2)
    phi_c = math.radians(lat_c_deg)
    lam_c = math.radians(lon_c_deg)
    alpha_c = math.radians(azimuth_deg)
    gamma_c = math.radians(gamma_c_deg)

    B = math.sqrt(1 + e2 * math.cos(phi_c)**4 / (1 - e2))
    A = a * B * k0 * math.sqrt(1 - e2) / (1 - e2 * math.sin(phi_c)**2)

    t0 = math.tan(math.pi/4 - phi_c/2) / ((1 - e*math.sin(phi_c))/(1 + e*math.sin(phi_c)))**(e/2)
    D  = B * math.sqrt(1 - e2) / (math.cos(phi_c) * math.sqrt(1 - e2*math.sin(phi_c)**2))
    D2 = max(D*D, 1.0)
    F  = D + math.copysign(math.sqrt(D2 - 1), lat_c_deg)
    H  = F * t0**B
    G  = (F - 1/F) / 2
    gamma_0 = math.asin(math.sin(alpha_c) / D)
    lam_0   = lam_c - math.asin(G * math.tan(gamma_0)) / B

    return {"a": a, "e": e, "e2": e2, "B": B, "A": A, "F": F, "H": H,
            "G": G, "gamma_0": gamma_0, "gamma_c": gamma_c, "lam_0": lam_0,
            "alpha_c": alpha_c}


def geo_to_rso(rso_zone, lat_deg, lon_deg):
    """Geographic (GDM datum) → RSO (Malaya). Returns (Northing, Easting)."""
    c  = rso_zone["_c"]
    phi = math.radians(lat_deg)
    lam = math.radians(lon_deg)
    e = c["e"]; B = c["B"]; A = c["A"]

    t = math.tan(math.pi/4 - phi/2) / ((1 - e*math.sin(phi))/(1 + e*math.sin(phi)))**(e/2)
    Q = c["H"] / t**B
    S = (Q - 1/Q) / 2
    T = (Q + 1/Q) / 2
    V = math.sin(B * (lam - c["lam_0"]))
    U = (-V * math.cos(c["gamma_0"]) + S * math.sin(c["gamma_0"])) / T
    v = A * math.log((1 - U) / (1 + U)) / (2*B)
    u = A * math.atan2(S * math.cos(c["gamma_0"]) + V * math.sin(c["gamma_0"]),
                       math.cos(B * (lam - c["lam_0"]))) / B

    E = v * math.cos(c["gamma_c"]) + u * math.sin(c["gamma_c"]) + rso_zone["FE"]
    N = u * math.cos(c["gamma_c"]) - v * math.sin(c["gamma_c"]) + rso_zone["FN"]
    return N, E


def rso_to_geo(rso_zone, N_grid, E_grid):
    """RSO (Malaya) → Geographic (GDM datum). Returns (lat_deg, lon_deg)."""
    c  = rso_zone["_c"]
    e = c["e"]; B = c["B"]; A = c["A"]

    v_ = (E_grid - rso_zone["FE"]) * math.cos(c["gamma_c"]) \
       - (N_grid - rso_zone["FN"]) * math.sin(c["gamma_c"])
    u_ = (N_grid - rso_zone["FN"]) * math.cos(c["gamma_c"]) \
       + (E_grid - rso_zone["FE"]) * math.sin(c["gamma_c"])

    Q_ = math.exp(-B * v_ / A)
    S_ = (Q_ - 1/Q_) / 2
    T_ = (Q_ + 1/Q_) / 2
    V_ = math.sin(B * u_ / A)
    U_ = (V_ * math.cos(c["gamma_0"]) + S_ * math.sin(c["gamma_0"])) / T_

    t_ = (c["H"] / math.sqrt((1 + U_) / (1 - U_)))**(1/B)
    # Iterative latitude from isometric latitude
    chi = math.pi/2 - 2*math.atan(t_)
    phi = chi
    for _ in range(15):
        phi_new = (math.pi/2
                   - 2*math.atan(t_ * ((1 - e*math.sin(phi))/(1 + e*math.sin(phi)))**(e/2)))
        if abs(phi_new - phi) < 1e-12:
            phi = phi_new; break
        phi = phi_new

    lam = c["lam_0"] - math.atan2(S_ * math.cos(c["gamma_0"]) - V_ * math.sin(c["gamma_0"]),
                                    math.cos(B * u_ / A)) / B
    return math.degrees(phi), math.degrees(lam)


def _init_zones_from_xml():
    """Load ZONES / RSO_ZONES / DATUM_TRANSFORMS from Malaysia.utm / Malaysia.xml."""
    global ZONES, RSO_ZONES, DATUM_TRANSFORMS, ZONE_NAMES
    _bin = os.path.dirname(os.path.abspath(__file__))
    if _bin not in sys.path:
        sys.path.insert(0, _bin)
    from app_paths import project_root
    from malaysia_zones import load_malaysia_config
    ZONES, RSO_ZONES, DATUM_TRANSFORMS = load_malaysia_config(
        project_root(), MOD_EVEREST, GRS80, _rso_constants)
    ZONE_NAMES = [z["name"] for z in ZONES]
    set_transforms(DATUM_TRANSFORMS)


_init_zones_from_xml()


# ─────────────────────────────────────────────────────────────────────────────
#  KML / KMZ helpers
# ─────────────────────────────────────────────────────────────────────────────
def _build_kml(points):
    """points: list of (name, lat_dd, lon_dd, description)"""
    rows = []
    for name, lat, lon, desc in points:
        rows.append(f"""  <Placemark>
    <name>{name}</name>
    <description>{desc}</description>
    <Point><coordinates>{lon:.8f},{lat:.8f},0</coordinates></Point>
  </Placemark>""")
    body = "\n".join(rows)
    return (f'<?xml version="1.0" encoding="UTF-8"?>\n'
            f'<kml xmlns="http://www.opengis.net/kml/2.2">\n'
            f'<Document>\n  <name>{app_name()} Export</name>\n'
            f'{body}\n</Document>\n</kml>')

def save_kml(filepath, points):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(_build_kml(points))

def save_kmz(filepath, points):
    kml = _build_kml(points)
    with zipfile.ZipFile(filepath, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("doc.kml", kml)


# ─────────────────────────────────────────────────────────────────────────────
#  GUI CONSTANTS
# ─────────────────────────────────────────────────────────────────────────────
DARK_BG  = "#1B2838"
PANEL_BG = "#243447"
CARD_BG  = "#1E3250"
ACCENT   = "#4FC3F7"
ACCENT2  = "#66BB6A"
TEXT_MAIN= "#E8F0FE"
TEXT_DIM = "#90A4AE"
INPUT_BG = "#152238"
INPUT_FG = "#E3F2FD"
ERROR_FG = "#EF9A9A"
BTN_BG   = "#1565C0"
BTN_HOV  = "#1976D2"
FONT_LBL = ("Segoe UI", 10)
FONT_HDR = ("Segoe UI", 11, "bold")
FONT_SM  = ("Segoe UI", 9)
FONT_MONO= ("Consolas", 11)
FONT_BIG = ("Consolas", 13, "bold")
FONT_TTL = ("Segoe UI", 14, "bold")


# ─────────────────────────────────────────────────────────────────────────────
#  WIDGETS
# ─────────────────────────────────────────────────────────────────────────────
class DMSEntry(tk.Frame):
    def __init__(self, parent, hemispheres=("N","S"), **kw):
        super().__init__(parent, bg=PANEL_BG, **kw)
        self.hemispheres = hemispheres
        self.d_var = tk.StringVar()
        self.m_var = tk.StringVar()
        self.s_var = tk.StringVar()
        self.h_var = tk.StringVar(value=hemispheres[0])
        root = parent.winfo_toplevel()
        if hasattr(root, "_entry_cfg"):
            ecfg = root._entry_cfg()
            lcfg = dict(bg=root._t["panel"], fg=root._t["text_dim"], font=FONT_SM)
        else:
            ecfg = dict(bg=INPUT_BG, fg=INPUT_FG, font=FONT_MONO,
                        insertbackground=ACCENT, relief="flat",
                        highlightthickness=1, highlightbackground="#37474F",
                        highlightcolor=ACCENT)
            lcfg = dict(bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.e_d = tk.Entry(self, textvariable=self.d_var, width=5, **ecfg)
        tk.Label(self, text="°", **lcfg).grid(row=0, column=1, padx=(1,6))
        self.e_m = tk.Entry(self, textvariable=self.m_var, width=4, **ecfg)
        tk.Label(self, text="′", **lcfg).grid(row=0, column=3, padx=(1,6))
        self.e_s = tk.Entry(self, textvariable=self.s_var, width=12, **ecfg)
        tk.Label(self, text="″", **lcfg).grid(row=0, column=5, padx=(1,6))
        self.cb = ttk.Combobox(self, textvariable=self.h_var,
                                values=list(hemispheres), width=3,
                                state="readonly", font=FONT_SM)
        self.e_d.grid(row=0, column=0)
        self.e_m.grid(row=0, column=2)
        self.e_s.grid(row=0, column=4)
        self.cb.grid(row=0, column=6)

    def get_dd(self):
        return dms_to_dd(float(self.d_var.get() or 0),
                         float(self.m_var.get() or 0),
                         float(self.s_var.get() or 0),
                         self.h_var.get())

    def set_dd(self, dd):
        d, m, s = dd_to_dms(dd)
        self.d_var.set(str(d)); self.m_var.set(str(m)); self.s_var.set(f"{s:.{_DMS_SEC_DECIMALS}f}")
        self.h_var.set(self.hemispheres[0] if dd >= 0 else self.hemispheres[1])

    def clear(self):
        for v in (self.d_var, self.m_var, self.s_var): v.set("")
        self.h_var.set(self.hemispheres[0])


class ResultCard(tk.Frame):
    def __init__(self, parent, label, **kw):
        super().__init__(parent, bg=CARD_BG, padx=12, pady=8, **kw)
        self._hdr = tk.Label(self, text=label, bg=CARD_BG, fg=TEXT_DIM, font=FONT_SM)
        self._hdr.pack(anchor="w")
        self.var = tk.StringVar(value="—")
        self._val_lbl = tk.Label(self, textvariable=self.var, bg=CARD_BG, fg=ACCENT2,
                 font=FONT_BIG, anchor="w")
        self._val_lbl.pack(anchor="w", fill="x")
        root = self.winfo_toplevel()
        if hasattr(root, "_result_cards"):
            root._result_cards.append(self)

    def set(self, v): self.var.set(v)
    def clear(self): self.var.set("—")
    def set_label(self, text): self._hdr.config(text=text)


# ─────────────────────────────────────────────────────────────────────────────
#  UI strings (PUBLIC_BUILD strips parameter / method detail)
# ─────────────────────────────────────────────────────────────────────────────
UTM_APP_NAME = "UTM Coordinate Wizard"


def app_name() -> str:
    return UTM_APP_NAME


def footer_text() -> str:
    return (
        "Copyright © 2026 Created and Modified by Firdaus Shah  |  "
        "From UTM Student, For UTM Student."
    )


def _help_text():
    name = app_name()
    lines = [
        name,
        "═" * len(name),
        "",
        "How to use each tab — follow the steps in order.",
        "",
        "── CONVERT TAB ──────────────────────────────────────────────",
        "",
        "  Convert one point between coordinate types.",
        "",
        "  ★ FASTEST — paste a coordinate pair into Input",
        "     Copy two values from Excel, Notepad, or a report. Click the field",
        "     that matches your data order and press  Ctrl+V  (or type both values",
        "     in one field separated by comma, semicolon, tab, or space).",
        "",
        "     Which field to paste into:",
        "       • Lat / Lon — paste into either field; the program auto-detects",
        "         which value is latitude (|value| ≤ 90°) and which is longitude.",
        "       • Northing / Easting (or y / x) — order depends on the field:",
        "           – y first  → paste into  Northing",
        "           – x first  → paste into  Easting",
        "     Both fields fill from one paste — then click  Convert .",
        "",
        "  Step by step:",
        "  1. Choose  From  (what you have) and  To  (what you need) at the top.",
        "  2. Use the datum, state/zone, and option rows that appear for your choices.",
        "  3. Enter or paste coordinates in the Input panel:",
        "       • Geographic — decimal degrees or DMS (toggle top-right of Input).",
        "       • Cassini / RSO / UTM — Northing and Easting (zone for UTM).",
        "       • MRT48 Cassini — pick Metric or Links to match your plan.",
        "  4. Click  Convert .  Results appear in the Output panel.",
        "  5. Optional:  Export KML/KMZ  to open the point in Google Earth.",
        "  6. Click  Clear  to reset and convert another point.",
        "",
        "── BATCH TAB ────────────────────────────────────────────────",
        "",
        "  Convert many points at once.",
        "",
        "  ★ FASTEST — paste many rows from Excel or a spreadsheet",
        "     Select your coordinate columns in Excel (or similar),  Ctrl+C .",
        "     On Batch, click in the table and press  Ctrl+V  (or click",
        "     Paste (Ctrl+V) ). Each clipboard row becomes one table row.",
        "     Column order:",
        "       • Lat/Lon input — auto-detected per row (|value| ≤ 90° = latitude).",
        "       • Northing/Easting — choose  Paste: N then E  or  E then N  above",
        "         the table to match your spreadsheet (or use a header row: Northing,",
        "         Easting, Lat, Lon, etc.).",
        "     A header row is skipped automatically when detected.",
        "     Coordinates-only paste (no name column) fills Northing/Easting or",
        "     Lat/Lon — the # column auto-numbers rows; Name stays blank.",
        "     When  From  is UTM, set Zone and Hemisphere above the table.",
        "",
        "  Other ways to fill the table:",
        "  1. Set  From  →  To  and the options row below (same choices as Convert).",
        "  2. Add data:",
        "       • Paste from Excel / tab-separated text (see above), or",
        "       • Type directly into the grid, or",
        "       •  Import TXT/CSV  or  Import Excel  (whole file; comma/semicolon OK).",
        "  3. Click  Convert All .  Each row shows OK or an error message.",
        "  4. Export:  Excel ,  CSV , or  KML/KMZ  (all valid points on a map).",
        "  5.  Clear All  empties the table for a new job.",
        "",
        "── GEOID HEIGHT TAB ─────────────────────────────────────────",
        "",
        "  Look up geoid undulation (N) and related heights.",
        "",
        "  1. Supply the position either way:",
        "       • Tick  Mirror projection and coordinates from Convert tab  to reuse",
        "         your Convert settings and coordinates, or",
        "       • Leave mirroring off — set  From  and enter coordinates on this tab",
        "         (same paste rules as Convert: lat/lon auto-detect; y→Northing,",
        "         x→Easting for grid coordinates).",
        "  2. Optional heights — enter ellipsoidal  h  and/or orthometric  H :",
        "       leave both blank for N only; fill one to calculate the other.",
        "  3. Click  Compute .  N and any derived heights appear in Results.",
        "  4. Click  Clear  to reset.",
    ]
    if LECTURER_MODE:
        lines += [
            "",
            "── DISTRIBUTE TAB (LECTURER) ──────────────────────────────────",
            "",
            "  Share the student edition with your class.",
            "",
            "  1. Optionally type your name (shown as From your lecturer name in the student README).",
            "  2. Click  Export Student ZIP…  and choose where to save.",
            "  3. Give students the zip — they run the .exe inside; no extra files needed.",
            "  4. Optional:  Generate class test coordinates…  for practice sets.",
        ]
    lines += [
        "",
        "── QUICK TIPS ───────────────────────────────────────────────",
        "",
        "  • Paste split: lat/lon auto-detects by range; grid coords use the field",
        "    you paste into — Northing if y is first, Easting if x is first.",
        "  • Batch paste: lat/lon auto-detects; grid coords use  N then E / E then N",
        "    above the table, or column headers (Northing, Easting, …).",
        "  • Comma-separated text files: use  Import TXT/CSV  instead of paste.",
        "  • MRT48 Cassini: Metric vs Links must match your plan/drawing.",
        "  • Pahang: choose the MRT48 sub-zone for your survey area.",
        "  • Convert may warn if a result falls outside Peninsular Malaysia.",
    ]
    return "\n".join(lines) + "\n"


# ─────────────────────────────────────────────────────────────────────────────
#  MAIN APP
# ─────────────────────────────────────────────────────────────────────────────
class CassiniApp(tk.Tk):
    def __init__(self):
        from ui_theme import ensure_win_app_user_model_id
        ensure_win_app_user_model_id()
        super().__init__()
        self._theme_mode = load_theme_mode()
        self._t = dict(THEMES[self._theme_mode])
        self._result_cards = []
        titles = {
            "lecturer": f"{UTM_APP_NAME} — Lecturer Edition",
            "student": f"{UTM_APP_NAME} — Student Edition",
            "dev": f"{UTM_APP_NAME} — Dev",
        }
        self.title(titles.get(edition(), titles["dev"]))
        self.resizable(True, True)
        self.minsize(820, 580)
        self.configure(bg=self._t["window"])
        self._styles()
        self._build()
        self._apply_theme(self._theme_mode)
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        self.geometry(f"{w}x{h}+{(self.winfo_screenwidth()-w)//2}+{(self.winfo_screenheight()-h)//2}")

    def _styles(self):
        t = self._t
        s = ttk.Style(); s.theme_use("clam")
        s.configure("TNotebook", background=t["window"], borderwidth=0)
        s.configure("TNotebook.Tab", background=t["panel"], foreground=t["text_dim"],
                     padding=[12, 6], font=FONT_LBL)
        s.map("TNotebook.Tab",
              background=[("selected", t["card"])],
              foreground=[("selected", t["accent"])])
        for style in ("Zone.TCombobox", "H.TCombobox"):
            s.configure(style, fieldbackground=t["input_bg"], background=t["input_bg"],
                        foreground=t["input_fg"], selectbackground=t["input_bg"],
                        selectforeground=t["input_fg"])
            s.map(style,
                  fieldbackground=[("readonly", t["input_bg"]), ("disabled", t["input_bg"])],
                  foreground=[("readonly", t["input_fg"]), ("disabled", t["text_dim"])],
                  selectbackground=[("readonly", t["input_bg"])],
                  selectforeground=[("readonly", t["input_fg"])])
        self._apply_batch_tree_style()

    def _apply_batch_tree_style(self):
        t = self._t
        s = ttk.Style()
        s.configure("Batch.Treeview", font=("Consolas", 9), rowheight=20,
                    background=t["input_bg"], fieldbackground=t["input_bg"],
                    foreground=t["input_fg"])
        s.configure("Batch.Treeview.Heading", font=("Segoe UI", 9, "bold"),
                    background=t["card"], foreground=t["accent"])
        s.map("Batch.Treeview", background=[("selected", t["tree_sel"])],
              foreground=[("selected", "white")])

    def _toggle_theme(self):
        nxt = "light" if self._theme_mode == "dark" else "dark"
        self._apply_theme(nxt)

    def _apply_theme(self, mode: str):
        if mode not in THEMES:
            mode = "dark"
        self._theme_mode = mode
        self._t = dict(THEMES[mode])
        save_theme_mode(mode)
        t = self._t

        self.configure(bg=t["window"])
        self._styles()

        skip: set = set()
        for chrome in (getattr(self, "_hdr", None), getattr(self, "_footer", None)):
            if chrome:
                _collect_skip(chrome, skip)

        apply_theme_widgets(self, t, frozenset(skip))

        for frame in (getattr(self, "_hdr", None), getattr(self, "_title_frame", None),
                      getattr(self, "_footer", None)):
            if frame:
                frame.config(bg=t["hdr"])
        if getattr(self, "_hdr", None):
            for child in self._hdr.winfo_children():
                try:
                    if child is getattr(self, "_hdr_logo_lbl", None):
                        child.config(bg=t["hdr"])
                    else:
                        child.config(bg=t["hdr"])
                except tk.TclError:
                    pass
            if getattr(self, "_title_frame", None):
                self._title_frame.config(bg=t["hdr"])
        if getattr(self, "_hdr_title_lbl", None):
            self._hdr_title_lbl.config(bg=t["hdr"], fg=t["hdr_title"], text=app_name())
        if getattr(self, "_hdr_sub_lbl", None):
            self._hdr_sub_lbl.config(bg=t["hdr"], fg=t["hdr_sub"])
        if getattr(self, "_footer_lbl", None):
            self._footer_lbl.config(bg=t["hdr"], fg=t["hdr_foot"], text=footer_text())
        if getattr(self, "_theme_btn", None):
            self._theme_btn.config(
                bg=t["chrome_btn"], fg=t["hdr_title"],
                text="Dark mode" if mode == "light" else "Light mode",
            )
        if getattr(self, "_hdr_logo_lbl", None):
            self._hdr_logo_lbl.config(bg=t["hdr"])
        self._refresh_header_logo()

        if getattr(self, "batch_tree", None):
            self.batch_tree.tag_configure("ok", foreground=t["accent2"])
            self.batch_tree.tag_configure("err", foreground=t["error"])
        self._apply_result_cards()

    def _refresh_header_logo(self):
        if not getattr(self, "_hdr_logo_lbl", None):
            return
        import os as _os
        from app_paths import resource as _res
        from PIL import ImageTk
        t = self._t
        logo_path = _res("logo.png")
        try:
            img = header_logo_image(t["logo_bg"], logo_path)
            if img is None:
                return
            self._logo_img = ImageTk.PhotoImage(img)
            self._hdr_logo_lbl.config(image=self._logo_img, bg=t["hdr"])
            self._hdr_logo_lbl.image = self._logo_img
        except Exception:
            pass

    def _apply_result_cards(self):
        """Style ResultCard output panels with theme maroon/sand colours."""
        t = self._t
        out_bg = t.get("output_bg", t["card"])
        out_fg = t.get("output_fg", t["accent2"])
        for card in getattr(self, "_result_cards", []):
            try:
                card.config(bg=out_bg)
                for child in card.winfo_children():
                    if isinstance(child, tk.Label):
                        kw = {"bg": out_bg}
                        if child.cget("textvariable"):
                            kw["fg"] = out_fg
                        else:
                            kw["fg"] = t["text_dim"]
                        child.config(**kw)
            except tk.TclError:
                pass

    def _entry_cfg(self, width=None, font=None):
        t = self._t
        cfg = dict(
            bg=t["input_bg"], fg=t["input_fg"], font=font or FONT_MONO,
            insertbackground=t["accent"], relief="flat",
            highlightthickness=1, highlightbackground=t["entry_hl"],
            highlightcolor=t["accent"],
        )
        if width is not None:
            cfg["width"] = width
        return cfg

    def _bind_coord_pair(self, entry_a, entry_b, var_a, var_b, mode="grid"):
        """Auto-split paired coords.

        mode='grid'  — top field y,x; bottom field x,y (N/E).
        mode='latlon' — smart |lat|≤90 detection; field hint if ambiguous.
        """
        def apply_pair(pair, from_above):
            if mode == "latlon":
                err = validate_lat_lon_paste(pair)
                if err:
                    messagebox.showerror("Not Lat/Lon", err)
                    return
                lat, lon = assign_lat_lon(pair, from_above)
                var_a.set(lat)
                var_b.set(lon)
            elif from_above:
                var_a.set(pair[0])
                var_b.set(pair[1])
            else:
                var_b.set(pair[0])
                var_a.set(pair[1])

        def split_var(var, from_above):
            pair = split_coord_pair(var.get())
            if pair:
                apply_pair(pair, from_above)
                return True
            return False

        def on_focus_out(_event=None):
            if split_var(var_a, True):
                return
            split_var(var_b, False)

        def on_key_release(event):
            if event.char in (",", ";", "\t"):
                from_above = event.widget is entry_a
                split_var(var_a if from_above else var_b, from_above)

        def on_paste(from_above):
            def handler(_event):
                try:
                    clip = self.clipboard_get()
                except tk.TclError:
                    return
                pair = split_coord_pair(clip)
                if pair:
                    apply_pair(pair, from_above)
                    return "break"
            return handler

        entry_a.bind("<FocusOut>", on_focus_out, add="+")
        entry_b.bind("<FocusOut>", on_focus_out, add="+")
        entry_a.bind("<KeyRelease>", on_key_release, add="+")
        entry_b.bind("<KeyRelease>", on_key_release, add="+")
        entry_a.bind("<Control-v>", on_paste(True), add="+")
        entry_a.bind("<Control-V>", on_paste(True), add="+")
        entry_b.bind("<Control-v>", on_paste(False), add="+")
        entry_b.bind("<Control-V>", on_paste(False), add="+")

    def _btn(self, parent, text, cmd, accent=True):
        t = self._t
        bg = t["btn"] if accent else t["btn_sec"]
        hov = t["btn_hov"] if accent else t["btn_sec_hov"]
        b = tk.Button(parent, text=text, command=cmd, bg=bg, fg="white",
                      font=FONT_HDR, relief="flat", padx=14, pady=6, cursor="hand2",
                      activebackground=hov, activeforeground="white", bd=0)
        b.bind("<Enter>", lambda e: b.config(bg=hov))
        b.bind("<Leave>", lambda e: b.config(bg=bg))
        return b

    def _sec(self, p, t, fg=None):
        th = self._t
        tk.Label(p, text=t, bg=th["panel"], fg=fg or th["accent"],
                 font=FONT_HDR).pack(anchor="w", padx=16, pady=(10, 2))

    def _sep(self, p):
        tk.Frame(p, bg=self._t["sep"], height=1).pack(fill="x", padx=16, pady=8)

    def _scrollable_tab(self, parent):
        """Wrap tab content in a vertical scroll area; returns inner frame for widgets."""
        bg = self._t["panel"]
        outer = tk.Frame(parent, bg=bg)
        outer.pack(fill="both", expand=True)
        canvas = tk.Canvas(outer, bg=bg, highlightthickness=0, borderwidth=0)
        sb = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        inner = tk.Frame(canvas, bg=bg)
        cwin = canvas.create_window((0, 0), window=inner, anchor="nw")
        _BOUND = "_scroll_wheel_bound"

        def _resize(event):
            canvas.itemconfig(cwin, width=event.width)

        def _scrollregion(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _wheel(event):
            if getattr(event, "delta", 0):
                canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
            elif getattr(event, "num", None) == 4:
                canvas.yview_scroll(-1, "units")
            elif getattr(event, "num", None) == 5:
                canvas.yview_scroll(1, "units")
            return "break"

        def _bind_widget(widget):
            if getattr(widget, _BOUND, False):
                return
            widget.bind("<MouseWheel>", _wheel)
            widget.bind("<Button-4>", _wheel)
            widget.bind("<Button-5>", _wheel)
            setattr(widget, _BOUND, True)
            for child in widget.winfo_children():
                _bind_widget(child)

        def _bind_all_scroll_targets():
            _bind_widget(canvas)
            _bind_widget(outer)
            _bind_widget(sb)
            _bind_widget(inner)

        def _on_inner_configure(_event=None):
            _scrollregion()
            canvas.after_idle(_bind_all_scroll_targets)

        canvas.bind("<Configure>", _resize)
        inner.bind("<Configure>", _on_inner_configure)
        inner._refresh_scroll_wheel = _bind_all_scroll_targets
        _on_inner_configure()
        return inner

    def _finish_scrollable_tab(self, inner):
        refresh = getattr(inner, "_refresh_scroll_wheel", None)
        if refresh:
            self.after_idle(refresh)

    def _zone_row(self, parent, zone_attr):
        """Build a projection zone selector row inside a tab.
        Stores the zone StringVar as self.<zone_attr>."""
        zfr = tk.Frame(parent, bg=PANEL_BG, padx=16, pady=8)
        zfr.pack(fill="x")
        tk.Label(zfr, text="Projection Zone :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).grid(row=0, column=0, sticky="w", padx=(0,8))
        zone_var = tk.StringVar(value=ZONE_NAMES[0])
        ttk.Combobox(zfr, textvariable=zone_var, values=ZONE_NAMES,
                     width=46, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).grid(row=0, column=1, sticky="w")
        info = tk.Label(zfr, text="", bg=PANEL_BG, fg=ACCENT, font=FONT_SM)
        info.grid(row=1, column=0, columnspan=2, sticky="w", pady=(2,0))

        def _upd(*_):
            z = next(x for x in ZONES if x["name"] == zone_var.get())
            info.config(text=_zone_datum_blurb(z))

        zone_var.trace_add("write", _upd)
        _upd()
        setattr(self, zone_attr, zone_var)
        return zone_var

    # ── main build ─────────────────────────────────────────────────────────
    def _build(self):
        _bin = _os.path.dirname(_os.path.abspath(__file__))

        # Window icon — title bar + Windows taskbar (matches embedded .exe icon)
        from app_paths import resource as _res
        from ui_theme import set_tk_window_icon
        set_tk_window_icon(self)

        # Header bar — logo left, title right
        t = self._t
        self._hdr = tk.Frame(self, bg=t["hdr"], height=64)
        self._hdr.pack(fill="x")
        self._hdr.pack_propagate(False)

        # Logo (left)
        from PIL import ImageTk
        logo_path = _res("logo.png")
        self._hdr_logo_lbl = tk.Label(self._hdr, bg=t["hdr"])
        try:
            img = header_logo_image(t["logo_bg"], logo_path if _os.path.exists(logo_path) else None)
            if img is not None:
                self._logo_img = ImageTk.PhotoImage(img)
                self._hdr_logo_lbl.config(image=self._logo_img)
                self._hdr_logo_lbl.image = self._logo_img
        except Exception as e:
            tk.Label(self._hdr, text=f"[logo error: {e}]",
                     font=("Arial", 7), bg=t["hdr"], fg=t["hdr_foot"]
                     ).pack(side="left", padx=16)
        self._hdr_logo_lbl.pack(side="left", padx=(6, 0), pady=4)

        # Theme toggle + title (right)
        right = tk.Frame(self._hdr, bg=t["hdr"])
        right.pack(side="right", padx=12)
        self._theme_btn = tk.Button(
            right, text="Light mode" if self._theme_mode == "dark" else "Dark mode",
            command=self._toggle_theme, bg=t["chrome_btn"], fg=t["hdr_title"],
            font=("Segoe UI", 9), relief="flat", padx=10, pady=4, cursor="hand2", bd=0)
        self._theme_btn.pack(side="right", padx=(8, 0))

        self._title_frame = tk.Frame(right, bg=t["hdr"])
        self._title_frame.pack(side="right", padx=8)
        self._hdr_title_lbl = tk.Label(self._title_frame, text=app_name(),
                                       font=("Arial", 16, "bold"), bg=t["hdr"], fg=t["hdr_title"])
        self._hdr_title_lbl.pack(anchor="e")
        if LECTURER_MODE:
            self._hdr_sub_lbl = tk.Label(
                self._title_frame, text="Lecturer Edition — classroom distribution enabled",
                font=("Arial", 9), bg=t["hdr"], fg=t["hdr_sub"])
            self._hdr_sub_lbl.pack(anchor="e")
        elif is_student():
            self._hdr_sub_lbl = tk.Label(
                self._title_frame, text="Student Edition",
                font=("Arial", 9), bg=t["hdr"], fg=t["hdr_sub"])
            self._hdr_sub_lbl.pack(anchor="e")
        elif not PUBLIC_BUILD:
            self._hdr_sub_lbl = tk.Label(
                self._title_frame, text="Source: Jabatan Ukur dan Pemetaan Malaysia (JUPEM)",
                font=("Arial", 9), bg=t["hdr"], fg=t["hdr_foot"])
            self._hdr_sub_lbl.pack(anchor="e")
        else:
            self._hdr_sub_lbl = None

        # Copyright footer — packed before notebook so it stays at bottom
        self._footer = tk.Frame(self, bg=t["hdr"], height=22)
        self._footer.pack(fill="x", side="bottom")
        self._footer.pack_propagate(False)
        self._footer_lbl = tk.Label(
            self._footer,
            text=footer_text(),
            font=("Arial", 8), bg=t["hdr"], fg=t["hdr_foot"])
        self._footer_lbl.pack(expand=True)

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=16, pady=(0,14))
        t1=tk.Frame(nb,bg=PANEL_BG); t2=tk.Frame(nb,bg=PANEL_BG)
        t3=tk.Frame(nb,bg=PANEL_BG); t4=tk.Frame(nb,bg=PANEL_BG)
        nb.add(t1, text="  Convert  ")
        nb.add(t2, text="  Batch  ")
        nb.add(t3, text="  Geoid Height  ")
        if LECTURER_MODE:
            t5 = tk.Frame(nb, bg=PANEL_BG)
            nb.add(t5, text="  Distribute  ")
            self._tab_distribute(t5)
        nb.add(t4, text="  Help  ")
        self._tab_convert(self._scrollable_tab(t1))
        self._tab_batch_kml(t2)
        self._tab_geoid(self._scrollable_tab(t3))
        self._tab_help(t4)

    # ── Convert tab ────────────────────────────────────────────────────────
    _CV_SYS = ["Geographic (Lat/Lon)", "Cassini-Soldner", "RSO", "UTM"]
    _CAS = "Cassini-Soldner"

    @staticmethod
    def _zone_label(name):
        return name.split("  (")[0].strip()

    def _zones_for(self, datum):
        return [self._zone_label(z["name"]) for z in ZONES if z["datum"] == datum]

    def _find_zone(self, datum, label):
        for z in ZONES:
            if z["datum"] == datum and self._zone_label(z["name"]) == label:
                return z
        raise ValueError(f"Zone '{label}' not found for datum {datum}")

    def _cv_rb_kwargs(self, command=None):
        return dict(bg=PANEL_BG, fg=TEXT_MAIN, selectcolor=INPUT_BG,
                    activebackground=PANEL_BG, activeforeground=ACCENT,
                    font=FONT_SM, command=command or self._cv_update)

    def _cv_make_cas_family_fr(self, parent, var, on_change=None):
        fr = tk.Frame(parent, bg=PANEL_BG)
        tk.Label(fr, text="Cassini datum :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        rb = self._cv_rb_kwargs(on_change)
        for txt, val in [("GDM2000", "GDM"), ("MRT48", "MRT48")]:
            tk.Radiobutton(fr, text=txt, variable=var, value=val,
                           **rb).pack(side="left", padx=6)
        return fr

    def _cv_make_geo_datum_fr(self, parent, var, on_change=None):
        fr = tk.Frame(parent, bg=PANEL_BG)
        tk.Label(fr, text="Geo datum :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        rb = self._cv_rb_kwargs(on_change)
        for txt, val in [("WGS84", "WGS84"), ("PMGSN94", "PMGSN94"), ("GDM2000", "GDM"), ("MRT48", "MRT48")]:
            tk.Radiobutton(fr, text=txt, variable=var, value=val,
                           **rb).pack(side="left", padx=6)
        return fr

    def _cv_make_gdm_rev_fr(self, parent, var, on_change=None):
        fr = tk.Frame(parent, bg=PANEL_BG)
        tk.Label(fr, text="Revision :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        rb = self._cv_rb_kwargs(on_change)
        for txt, val in [("2000", "GDM2000"), ("2009", "GDM2009")]:
            tk.Radiobutton(fr, text=txt, variable=var, value=val,
                           **rb).pack(side="left", padx=6)
        return fr

    def _cv_make_rso_family_fr(self, parent, var=None, on_change=None):
        var = var if var is not None else self.cv_rso_family
        fr = tk.Frame(parent, bg=PANEL_BG)
        tk.Label(fr, text="RSO datum :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        rb = self._cv_rb_kwargs(on_change)
        for txt, val in [("GDM2000", "GDM"), ("MRT48", "MRT48")]:
            tk.Radiobutton(fr, text=txt, variable=var, value=val,
                           **rb).pack(side="left", padx=6)
        return fr

    def _cv_rb_kwargs_geoid(self):
        return dict(bg=PANEL_BG, fg=TEXT_MAIN, selectcolor=INPUT_BG,
                    activebackground=PANEL_BG, activeforeground=ACCENT,
                    font=FONT_SM, command=self._geoid_local_update)

    def _g_geo_datum(self):
        fam = self.g_from_geo_family.get()
        if fam == "WGS84":
            return "WGS84"
        if fam == "GDM":
            return self.g_from_gdm_rev.get()
        if fam == "MRT48":
            return "MRT48"
        return "PMGSN94"

    def _g_cassini_datum(self):
        fam = self.g_from_cas_family.get()
        if fam == "GDM":
            return self.g_from_gdm_rev.get()
        return "MRT48"

    def _g_rso_zone_resolved(self):
        if self.g_rso_family.get() == "MRT48":
            return next(z for z in RSO_ZONES if z["datum"] == "MRT48")
        return next(z for z in RSO_ZONES if z["datum"] == self.g_from_gdm_rev.get())

    def _g_sys_datum(self, sys_name):
        if sys_name == "Geographic (Lat/Lon)":
            return self._g_geo_datum()
        if sys_name == self._CAS:
            return self._g_cassini_datum()
        if sys_name == "RSO":
            return ("MRT48" if self.g_rso_family.get() == "MRT48"
                    else self.g_from_gdm_rev.get())
        if sys_name == "UTM":
            return "WGS84"
        return self._g_cassini_datum()

    def _g_datum_blurb(self, sys_name):
        if sys_name == "Geographic (Lat/Lon)":
            d = self._g_geo_datum()
            ell = ("Modified Everest" if d == "MRT48"
                   else "WGS84" if d in ("WGS84", "PMGSN94") else "GRS80")
            return f"Datum: {d}   Ellipsoid: {ell}"
        if sys_name == "UTM":
            return "Datum: WGS84   Ellipsoid: WGS84   Standard UTM (k₀=0.9996)"
        if sys_name == "RSO":
            return _rso_datum_blurb(self._g_rso_zone_resolved())
        if sys_name == self._CAS:
            dkey = self._g_cassini_datum()
            zone_name = self.g_from_zone.get()
            if zone_name:
                try:
                    return _zone_datum_blurb(self._find_zone(dkey, zone_name))
                except ValueError:
                    pass
            ell = "Modified Everest" if dkey == "MRT48" else "GRS80"
            return f"Datum: {dkey}   Ellipsoid: {ell}"
        return ""

    def _geoid_get_latlon(self):
        if self.g_mode.get() == "dd":
            return float(self.g_lat_dd.get()), float(self.g_lon_dd.get())
        return self.g_lat_dms.get_dd(), self.g_lon_dms.get_dd()

    def _geoid_refresh_mirror_summary(self):
        if not getattr(self, "g_mirror_lbl", None):
            return
        if not self.geoid_mirror_convert.get():
            return
        frm = self.cv_from.get()
        zone = self.cv_from_zone.get() if frm == self._CAS else None
        blurb = self._cv_datum_blurb(frm, zone, "from")
        self.g_mirror_lbl.config(
            text=f"Using Convert → From:  {frm}\n{blurb}",
            fg=ACCENT)

    def _geoid_sync_toggle(self):
        mirror = self.geoid_mirror_convert.get()
        if mirror:
            self.g_local_panel.pack_forget()
            self.g_mirror_panel.pack(fill="x", padx=16, pady=(4, 0))
            self._geoid_refresh_mirror_summary()
        else:
            self.g_mirror_panel.pack_forget()
            self.g_local_panel.pack(fill="x", padx=0, pady=(4, 0))
            self._geoid_local_update()
        inner = getattr(self, "_geoid_scroll_inner", None)
        if inner:
            self._finish_scrollable_tab(inner)

    def _geoid_refresh_datum_info(self):
        if not getattr(self, "g_from_datum_info", None):
            return
        frm = self.g_from.get()
        from_zone = self.g_from_zone.get() if frm == self._CAS else None
        self.g_from_datum_info.config(text=self._g_datum_blurb(frm))

    def _geoid_layout_from_options(self, frm_cas, frm_rso, frm_geo, rso_gdm):
        for fr in (self.g_from_cas_fr, self.g_from_geo_fr,
                   self.g_from_gdm_rev_fr, self.g_from_rso_fr):
            fr.pack_forget()
        from_cas_gdm = frm_cas and self.g_from_cas_family.get() == "GDM"
        from_geo_gdm = frm_geo and self.g_from_geo_family.get() == "GDM"
        stack = []
        if frm_geo:
            stack.append(self.g_from_geo_fr)
        if frm_cas:
            stack.append(self.g_from_cas_fr)
        if frm_rso:
            stack.append(self.g_from_rso_fr)
        if from_cas_gdm or from_geo_gdm or (frm_rso and rso_gdm):
            stack.append(self.g_from_gdm_rev_fr)
        for fr in stack:
            fr.pack(anchor="w", pady=(4, 0))

    def _geoid_local_update(self, *_):
        if self.geoid_mirror_convert.get():
            return
        frm = self.g_from.get()
        frm_cas = frm == self._CAS
        frm_utm = frm == "UTM"
        frm_geo = frm == "Geographic (Lat/Lon)"
        frm_rso = frm == "RSO"
        rso_gdm = self.g_rso_family.get() == "GDM"

        self.g_from_state_lbl.pack_forget()
        self.g_from_zone_cb.pack_forget()
        if frm_cas:
            dkey = self._g_cassini_datum()
            zones = self._zones_for(dkey)
            self.g_from_zone_cb["values"] = zones
            if self.g_from_zone.get() not in zones:
                self.g_from_zone.set(zones[0] if zones else "")
            self.g_from_state_lbl.pack(side="left", padx=(0, 6))
            self.g_from_zone_cb.pack(side="left")

        self._geoid_layout_from_options(frm_cas, frm_rso, frm_geo, rso_gdm)
        self._geoid_refresh_datum_info()

        if frm_geo:
            self.g_mode_fr.pack(anchor="w", pady=(4, 0))
        else:
            self.g_mode_fr.pack_forget()

        self.g_inp_geo.pack_forget()
        self.g_inp_cas.pack_forget()
        self.g_inp_utm.pack_forget()
        if frm_geo:
            self.g_inp_geo.pack(anchor="w", fill="x")
            if self.g_mode.get() == "dd":
                self.g_geo_dms_fr.pack_forget()
                self.g_geo_dd_fr.pack(anchor="w", fill="x")
            else:
                self.g_geo_dd_fr.pack_forget()
                self.g_geo_dms_fr.pack(anchor="w", fill="x")
        elif frm_cas or frm_rso:
            self.g_inp_cas.pack(anchor="w", fill="x")
        elif frm_utm:
            self.g_inp_utm.pack(anchor="w", fill="x")
        self._geoid_mrt48_unit_changed()

        if frm_utm:
            self.g_utm_note.pack(anchor="w", pady=(4, 0))
        else:
            self.g_utm_note.pack_forget()
        inner = getattr(self, "_geoid_scroll_inner", None)
        if inner:
            self._finish_scrollable_tab(inner)

    def _geoid_dms_toggle(self):
        if self.g_mode.get() == "dd":
            self.g_geo_dms_fr.pack_forget()
            self.g_geo_dd_fr.pack(anchor="w", fill="x")
        else:
            self.g_geo_dd_fr.pack_forget()
            self.g_geo_dms_fr.pack(anchor="w", fill="x")

    def _geoid_init_local_vars(self):
        """Independent Geoid-tab projection + coordinate variables."""
        self.g_from = tk.StringVar(value="Geographic (Lat/Lon)")
        self.g_from_zone = tk.StringVar()
        self.g_from_cas_family = tk.StringVar(value="MRT48")
        self.g_from_gdm_rev = tk.StringVar(value="GDM2009")
        self.g_from_geo_family = tk.StringVar(value="WGS84")
        self.g_rso_family = tk.StringVar(value="GDM")
        self.g_mode = tk.StringVar(value="dd")
        self.g_lat_dd = tk.StringVar()
        self.g_lon_dd = tk.StringVar()
        self.g_cas_N = tk.StringVar()
        self.g_cas_E = tk.StringVar()
        self.g_utm_N = tk.StringVar()
        self.g_utm_E = tk.StringVar()
        self.g_utm_zone = tk.StringVar()
        self.g_utm_hemi = tk.StringVar(value="N")
        self.g_mrt48_input_unit = tk.StringVar(value="m")

    def _geoid_build_input_section(self, p):
        """Projection + coordinates: mirror Convert tab or enter on Geoid tab."""
        self._geoid_init_local_vars()
        self.geoid_mirror_convert = tk.BooleanVar(value=False)

        self._sec(p, "Input projection & coordinates")
        sync_fr = tk.Frame(p, bg=PANEL_BG)
        sync_fr.pack(anchor="w", padx=16, pady=(4, 0))
        tk.Checkbutton(
            sync_fr,
            text="Mirror projection and coordinates from Convert tab",
            variable=self.geoid_mirror_convert,
            bg=PANEL_BG, fg=TEXT_MAIN, selectcolor=INPUT_BG,
            activebackground=PANEL_BG, activeforeground=ACCENT,
            font=FONT_LBL, command=self._geoid_sync_toggle,
        ).pack(anchor="w")

        self.g_mirror_panel = tk.Frame(p, bg=PANEL_BG)
        self.g_mirror_lbl = tk.Label(
            self.g_mirror_panel, text="", bg=PANEL_BG, fg=ACCENT,
            font=FONT_SM, justify="left", wraplength=620)
        self.g_mirror_lbl.pack(anchor="w", pady=(4, 0))
        tk.Label(
            self.g_mirror_panel,
            text="Change projection or coordinates on the Convert tab — they apply here while mirroring is on.",
            bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM, wraplength=620,
            justify="left",
        ).pack(anchor="w", pady=(2, 0))

        self.g_local_panel = tk.Frame(p, bg=PANEL_BG)
        ecfg = self._entry_cfg(width=22)
        proj = tk.Frame(self.g_local_panel, bg=PANEL_BG)
        proj.pack(fill="x", padx=16, pady=(4, 0))

        row0 = tk.Frame(proj, bg=PANEL_BG)
        row0.pack(anchor="w", fill="x")
        tk.Label(row0, text="From :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        ttk.Combobox(row0, textvariable=self.g_from, values=self._CV_SYS,
                     width=22, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).pack(side="left")
        self.g_from.trace_add("write", lambda *_: self._geoid_local_update())
        self.g_mrt48_input_unit.trace_add("write", lambda *_: self._geoid_mrt48_unit_changed())

        row1 = tk.Frame(proj, bg=PANEL_BG)
        row1.pack(anchor="w", fill="x", pady=(2, 0))
        self.g_from_state_lbl = tk.Label(row1, text="State :", bg=PANEL_BG,
                                         fg=TEXT_DIM, font=FONT_SM)
        self.g_from_zone_cb = ttk.Combobox(
            row1, textvariable=self.g_from_zone, width=20, state="readonly",
            style="Zone.TCombobox", font=FONT_SM)
        self.g_from_zone.trace_add("write", lambda *_: self._geoid_local_update())

        self.g_from_datum_info = tk.Label(proj, text="", bg=PANEL_BG, fg=ACCENT,
                                          font=FONT_SM, anchor="w")
        self.g_from_datum_info.pack(anchor="w", pady=(2, 0))

        opts = tk.Frame(proj, bg=PANEL_BG)
        opts.pack(anchor="w", fill="x")
        rb = self._cv_rb_kwargs_geoid()
        self.g_from_cas_fr = tk.Frame(opts, bg=PANEL_BG)
        tk.Label(self.g_from_cas_fr, text="Cassini datum :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in [("GDM2000", "GDM"), ("MRT48", "MRT48")]:
            tk.Radiobutton(self.g_from_cas_fr, text=txt, variable=self.g_from_cas_family,
                           value=val, **rb).pack(side="left", padx=6)
        self.g_from_geo_fr = tk.Frame(opts, bg=PANEL_BG)
        tk.Label(self.g_from_geo_fr, text="Geo datum :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in [("WGS84", "WGS84"), ("PMGSN94", "PMGSN94"), ("GDM2000", "GDM"), ("MRT48", "MRT48")]:
            tk.Radiobutton(self.g_from_geo_fr, text=txt, variable=self.g_from_geo_family,
                           value=val, **rb).pack(side="left", padx=6)
        self.g_from_gdm_rev_fr = tk.Frame(opts, bg=PANEL_BG)
        tk.Label(self.g_from_gdm_rev_fr, text="Revision :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in [("2000", "GDM2000"), ("2009", "GDM2009")]:
            tk.Radiobutton(self.g_from_gdm_rev_fr, text=txt, variable=self.g_from_gdm_rev,
                           value=val, **rb).pack(side="left", padx=6)
        self.g_from_rso_fr = self._cv_make_rso_family_fr(opts, self.g_rso_family)
        for w in self.g_from_rso_fr.winfo_children():
            if isinstance(w, tk.Radiobutton):
                w.config(command=self._geoid_local_update)
        self.g_utm_note = tk.Label(opts, text="UTM grid: WGS84 only",
                                   bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)

        hdr_fr = tk.Frame(self.g_local_panel, bg=PANEL_BG)
        hdr_fr.pack(fill="x", padx=16, pady=(8, 0))
        tk.Label(hdr_fr, text="Coordinates", bg=PANEL_BG, fg=ACCENT2,
                 font=FONT_HDR).pack(side="left")
        self.g_mode_fr = tk.Frame(hdr_fr, bg=PANEL_BG)
        for txt, val in [("Decimal Degrees", "dd"),
                         ("Degrees °   Minutes ′   Seconds ″", "dms")]:
            tk.Radiobutton(self.g_mode_fr, text=txt, variable=self.g_mode, value=val,
                           bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._geoid_dms_toggle).pack(side="left", padx=4)

        inp_host = tk.Frame(self.g_local_panel, bg=PANEL_BG)
        inp_host.pack(anchor="w", padx=16, fill="x")

        self.g_inp_geo = tk.Frame(inp_host, bg=PANEL_BG)
        self.g_geo_dd_fr = tk.Frame(self.g_inp_geo, bg=PANEL_BG)
        self.g_geo_dd_fr.pack(anchor="w", fill="x")
        self.g_geo_dms_fr = tk.Frame(self.g_inp_geo, bg=PANEL_BG)
        lat_e = lon_e = None
        for var, lbl in [(self.g_lat_dd, "Latitude   (decimal °)"),
                         (self.g_lon_dd, "Longitude  (decimal °)")]:
            r = tk.Frame(self.g_geo_dd_fr, bg=PANEL_BG)
            r.pack(anchor="w", pady=3)
            tk.Label(r, text=lbl, bg=PANEL_BG, fg=TEXT_MAIN,
                     font=FONT_LBL, width=22, anchor="w").pack(side="left")
            e = tk.Entry(r, textvariable=var, **ecfg)
            e.pack(side="left")
            if lbl.startswith("Latitude"):
                lat_e = e
            else:
                lon_e = e
        if lat_e and lon_e:
            self._bind_coord_pair(lat_e, lon_e, self.g_lat_dd, self.g_lon_dd, mode="latlon")
        lf = tk.Frame(self.g_geo_dms_fr, bg=PANEL_BG)
        lf.pack(anchor="w", pady=3)
        tk.Label(lf, text="Latitude", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=12, anchor="w").pack(side="left")
        self.g_lat_dms = DMSEntry(lf, hemispheres=("N", "S"))
        self.g_lat_dms.pack(side="left")
        lf2 = tk.Frame(self.g_geo_dms_fr, bg=PANEL_BG)
        lf2.pack(anchor="w", pady=3)
        tk.Label(lf2, text="Longitude", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=12, anchor="w").pack(side="left")
        self.g_lon_dms = DMSEntry(lf2, hemispheres=("E", "W"))
        self.g_lon_dms.pack(side="left")

        self.g_inp_cas = tk.Frame(inp_host, bg=PANEL_BG)
        self.g_mrt48_unit_fr = tk.Frame(self.g_inp_cas, bg=PANEL_BG)
        g_in_row = tk.Frame(self.g_mrt48_unit_fr, bg=PANEL_BG)
        g_in_row.pack(anchor="w")
        tk.Label(g_in_row, text="Input units :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in _MRT48_UNIT_OPTS:
            tk.Radiobutton(g_in_row, text=txt, variable=self.g_mrt48_input_unit,
                           value=val, bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._geoid_mrt48_unit_changed).pack(side="left", padx=4)
        self.g_mrt48_in_hint_lbl = tk.Label(
            g_in_row, text=mrt48_unit_format_hint("m"), bg=PANEL_BG,
            fg=TEXT_DIM, font=FONT_SM)
        self.g_mrt48_in_hint_lbl.pack(side="left", padx=(16, 0))
        self.g_mrt48_unit_fr.pack_forget()
        self.g_cas_N_lbl = tk.Label(self.g_inp_cas, bg=PANEL_BG)
        self.g_cas_E_lbl = tk.Label(self.g_inp_cas, bg=PANEL_BG)
        for var, lbl_attr, lbl_txt in [
                (self.g_cas_N, "g_cas_N_lbl", "Northing  [m]"),
                (self.g_cas_E, "g_cas_E_lbl", "Easting   [m]")]:
            r = tk.Frame(self.g_inp_cas, bg=PANEL_BG)
            r.pack(anchor="w", pady=3)
            lbl = getattr(self, lbl_attr)
            lbl.config(text=lbl_txt, fg=TEXT_MAIN, font=FONT_LBL, width=22, anchor="w")
            lbl.pack(side="left")
            tk.Entry(r, textvariable=var, **ecfg).pack(side="left")

        self.g_inp_utm = tk.Frame(inp_host, bg=PANEL_BG)
        for var, lbl in [(self.g_utm_N, "Northing  [m]"),
                         (self.g_utm_E, "Easting   [m]")]:
            r = tk.Frame(self.g_inp_utm, bg=PANEL_BG)
            r.pack(anchor="w", pady=3)
            tk.Label(r, text=lbl, bg=PANEL_BG, fg=TEXT_MAIN,
                     font=FONT_LBL, width=22, anchor="w").pack(side="left")
            tk.Entry(r, textvariable=var, **ecfg).pack(side="left")
        zfr = tk.Frame(self.g_inp_utm, bg=PANEL_BG)
        zfr.pack(anchor="w", pady=3)
        tk.Label(zfr, text="UTM Zone :", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        tk.Entry(zfr, textvariable=self.g_utm_zone,
                 **{**ecfg, "width": 6}).pack(side="left", padx=(0, 16))
        tk.Label(zfr, text="Hemisphere :", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL).pack(side="left")
        ttk.Combobox(zfr, textvariable=self.g_utm_hemi, values=["N", "S"],
                     width=4, state="readonly", style="H.TCombobox",
                     font=FONT_LBL).pack(side="left")

        self._geoid_sync_toggle()

    def _geoid_input_to_geographic(self):
        """Resolve Geoid-tab local input to (lat°, lon°, source_datum)."""
        frm = self.g_from.get()
        if frm == "Geographic (Lat/Lon)":
            lat, lon = self._geoid_get_latlon()
            return lat, lon, self._g_geo_datum()
        if frm == self._CAS:
            dkey = self._g_cassini_datum()
            z = self._find_zone(dkey, self.g_from_zone.get())
            if dkey == "MRT48":
                n, e = self._geoid_parse_mrt48_cassini(self.g_cas_N.get(), self.g_cas_E.get())
                lat, lon = mrt48_cassini_to_geo(z, n, e)
            else:
                lat, lon = cassini_to_geo(z, float(self.g_cas_N.get()), float(self.g_cas_E.get()))
            return lat, lon, dkey
        if frm == "UTM":
            zn = int(self.g_utm_zone.get())
            lat, lon = utm_to_geo(WGS84, float(self.g_utm_N.get()), float(self.g_utm_E.get()),
                                  zn, self.g_utm_hemi.get())
            return lat, lon, "WGS84"
        if frm == "RSO":
            rz = self._g_rso_zone_resolved()
            lat, lon = rso_to_geo(rz, float(self.g_cas_N.get()), float(self.g_cas_E.get()))
            return lat, lon, self._g_sys_datum("RSO")
        raise ValueError(f"Unknown input system: {frm}")

    def _tab_convert(self, p):
        ecfg = self._entry_cfg(width=22)

        # ── From / To selector row ──
        ft = tk.Frame(p, bg=PANEL_BG); ft.pack(fill="x", padx=16, pady=(10,4))

        # FROM column
        tk.Label(ft, text="From :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).grid(row=0, column=0, sticky="w", padx=(0,6))
        self.cv_from = tk.StringVar(value="Geographic (Lat/Lon)")
        ttk.Combobox(ft, textvariable=self.cv_from, values=self._CV_SYS,
                     width=22, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).grid(row=0, column=1, sticky="w")
        # From-State (Cassini)
        self.cv_from_state_lbl = tk.Label(ft, text="State :", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.cv_from_state_lbl.grid(row=1, column=0, sticky="w", padx=(0,6), pady=(2,0))
        self.cv_from_zone = tk.StringVar()
        self.cv_from_zone_cb = ttk.Combobox(ft, textvariable=self.cv_from_zone,
                                              width=20, state="readonly", style="Zone.TCombobox",
                                              font=FONT_SM)
        self.cv_from_zone_cb.grid(row=1, column=1, sticky="w", pady=(2,0))
        self.cv_from_state_lbl.grid_remove(); self.cv_from_zone_cb.grid_remove()
        self.cv_from_datum_info = tk.Label(ft, text="", bg=PANEL_BG, fg=ACCENT,
                                           font=FONT_SM, anchor="w")
        self.cv_from_datum_info.grid(row=2, column=0, columnspan=2, sticky="w", pady=(2, 0))

        # Arrow
        tk.Label(ft, text="→", bg=PANEL_BG, fg=ACCENT, font=("Segoe UI", 14, "bold")
                 ).grid(row=0, column=2, padx=18)

        # TO column
        tk.Label(ft, text="To :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).grid(row=0, column=3, sticky="w", padx=(0,6))
        self.cv_to = tk.StringVar(value=self._CAS)
        ttk.Combobox(ft, textvariable=self.cv_to, values=self._CV_SYS,
                     width=22, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).grid(row=0, column=4, sticky="w")
        # To-State (Cassini)
        self.cv_to_state_lbl = tk.Label(ft, text="State :", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.cv_to_state_lbl.grid(row=1, column=3, sticky="w", padx=(0,6), pady=(2,0))
        self.cv_to_zone = tk.StringVar()
        self.cv_to_zone_cb = ttk.Combobox(ft, textvariable=self.cv_to_zone,
                                            width=20, state="readonly", style="Zone.TCombobox",
                                            font=FONT_SM)
        self.cv_to_zone_cb.grid(row=1, column=4, sticky="w", pady=(2,0))
        # UTM zone override inline with To=UTM
        self.cv_to_utm_zone_lbl = tk.Label(ft, text="Zone (auto) :", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.cv_to_utm_zone_lbl.grid(row=1, column=3, sticky="w", padx=(0,6), pady=(2,0))
        self.cv_to_utm_override = tk.StringVar()
        self.cv_to_utm_zone_entry = tk.Entry(
            ft, textvariable=self.cv_to_utm_override, **self._entry_cfg(width=6))
        self.cv_to_utm_zone_entry.grid(row=1, column=4, sticky="w", pady=(2,0))
        self.cv_to_state_lbl.grid_remove(); self.cv_to_zone_cb.grid_remove()
        self.cv_to_utm_zone_lbl.grid_remove(); self.cv_to_utm_zone_entry.grid_remove()
        self.cv_to_datum_info = tk.Label(ft, text="", bg=PANEL_BG, fg=ACCENT,
                                         font=FONT_SM, anchor="w")
        self.cv_to_datum_info.grid(row=2, column=3, columnspan=2, sticky="w", pady=(2, 0))

        self.cv_from_cas_family = tk.StringVar(value="MRT48")
        self.cv_to_cas_family = tk.StringVar(value="MRT48")
        self.cv_from_gdm_rev = tk.StringVar(value="GDM2009")
        self.cv_to_gdm_rev = tk.StringVar(value="GDM2009")
        self.cv_from_geo_family = tk.StringVar(value="WGS84")
        self.cv_to_geo_family = tk.StringVar(value="WGS84")
        self.cv_rso_family = tk.StringVar(value="GDM")
        rb_kw = dict(bg=PANEL_BG, fg=TEXT_MAIN, selectcolor=INPUT_BG,
                     activebackground=PANEL_BG, activeforeground=ACCENT,
                     font=FONT_SM, command=self._cv_update)

        def _geo_datum_fr(var):
            fr = tk.Frame(ft, bg=PANEL_BG)
            tk.Label(fr, text="Geo datum :", bg=PANEL_BG, fg=TEXT_DIM,
                     font=FONT_SM).pack(side="left", padx=(0, 8))
            for txt, val in [("WGS84", "WGS84"), ("PMGSN94", "PMGSN94"), ("GDM2000", "GDM"), ("MRT48", "MRT48")]:
                tk.Radiobutton(fr, text=txt, variable=var, value=val,
                               **rb_kw).pack(side="left", padx=6)
            return fr

        def _cas_family_fr(var):
            fr = tk.Frame(ft, bg=PANEL_BG)
            tk.Label(fr, text="Cassini datum :", bg=PANEL_BG, fg=TEXT_DIM,
                     font=FONT_SM).pack(side="left", padx=(0, 8))
            for txt, val in [("GDM2000", "GDM"), ("MRT48", "MRT48")]:
                tk.Radiobutton(fr, text=txt, variable=var, value=val,
                               **rb_kw).pack(side="left", padx=6)
            return fr

        def _gdm_rev_fr(var):
            fr = tk.Frame(ft, bg=PANEL_BG)
            tk.Label(fr, text="Revision :", bg=PANEL_BG, fg=TEXT_DIM,
                     font=FONT_SM).pack(side="left", padx=(0, 8))
            for txt, val in [("2000", "GDM2000"), ("2009", "GDM2009")]:
                tk.Radiobutton(fr, text=txt, variable=var, value=val,
                               **rb_kw).pack(side="left", padx=6)
            return fr

        def _rso_fr():
            fr = tk.Frame(ft, bg=PANEL_BG)
            tk.Label(fr, text="RSO datum :", bg=PANEL_BG, fg=TEXT_DIM,
                     font=FONT_SM).pack(side="left", padx=(0, 8))
            for txt, val in [("GDM2000", "GDM"), ("MRT48", "MRT48")]:
                tk.Radiobutton(fr, text=txt, variable=self.cv_rso_family, value=val,
                               **rb_kw).pack(side="left", padx=6)
            return fr

        self.cv_from_cas_fr = _cas_family_fr(self.cv_from_cas_family)
        self.cv_to_cas_fr = _cas_family_fr(self.cv_to_cas_family)
        self.cv_from_geo_fr = _geo_datum_fr(self.cv_from_geo_family)
        self.cv_to_geo_fr = _geo_datum_fr(self.cv_to_geo_family)
        self.cv_from_gdm_rev_fr = _gdm_rev_fr(self.cv_from_gdm_rev)
        self.cv_to_gdm_rev_fr = _gdm_rev_fr(self.cv_to_gdm_rev)
        self.cv_from_rso_fr = _rso_fr()
        self.cv_to_rso_fr = _rso_fr()
        for fr in (self.cv_from_cas_fr, self.cv_to_cas_fr,
                   self.cv_from_geo_fr, self.cv_to_geo_fr,
                   self.cv_from_gdm_rev_fr, self.cv_to_gdm_rev_fr,
                   self.cv_from_rso_fr, self.cv_to_rso_fr):
            fr.grid_remove()

        # UTM is always WGS84 — no datum picker
        self.cv_utm_note = tk.Label(ft, text="UTM grid: WGS84 only", bg=PANEL_BG,
                                    fg=TEXT_DIM, font=FONT_SM)
        self.cv_utm_note.grid(row=4, column=0, columnspan=5, sticky="w", pady=(4, 0))
        self.cv_utm_note.grid_remove()

        self._sep(p)

        # ── Input section header + DD/DMS toggle on the right ──
        hdr_fr = tk.Frame(p, bg=PANEL_BG); hdr_fr.pack(fill="x", padx=16, pady=(4,0))
        self._sec(hdr_fr, "Input")
        self.cv_mode_fr = tk.Frame(hdr_fr, bg=PANEL_BG)
        self.cv_mode_fr.pack(side="right", padx=(0,4))
        self.cv_mode = tk.StringVar(value="dd")
        for txt, val in [("Decimal Degrees", "dd"),
                         ("Degrees °   Minutes ′   Seconds ″", "dms")]:
            tk.Radiobutton(self.cv_mode_fr, text=txt, variable=self.cv_mode, value=val,
                           bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._cv_dms_toggle).pack(side="left", padx=4)

        # ── Input panels ──
        inp_host = tk.Frame(p, bg=PANEL_BG); inp_host.pack(anchor="w", padx=16, fill="x")

        # Geo input (DD)
        self.cv_inp_geo = tk.Frame(inp_host, bg=PANEL_BG)
        self.cv_inp_geo.grid(row=0, column=0, sticky="w")
        geo_dd = tk.Frame(self.cv_inp_geo, bg=PANEL_BG)
        geo_dd.grid(row=0, column=0, sticky="w")
        geo_dms = tk.Frame(self.cv_inp_geo, bg=PANEL_BG)
        geo_dms.grid(row=0, column=0, sticky="w")
        geo_dms.grid_remove()
        self._cv_geo_dd_fr = geo_dd; self._cv_geo_dms_fr = geo_dms
        self.cv_lat_dd = tk.StringVar(); self.cv_lon_dd = tk.StringVar()
        r_lat = tk.Frame(geo_dd, bg=PANEL_BG); r_lat.pack(anchor="w", pady=3)
        tk.Label(r_lat, text="Latitude   (decimal °)", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        self.cv_lat_dd_entry = tk.Entry(r_lat, textvariable=self.cv_lat_dd, **ecfg)
        self.cv_lat_dd_entry.pack(side="left")
        r_lon = tk.Frame(geo_dd, bg=PANEL_BG); r_lon.pack(anchor="w", pady=3)
        tk.Label(r_lon, text="Longitude  (decimal °)", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        self.cv_lon_dd_entry = tk.Entry(r_lon, textvariable=self.cv_lon_dd, **ecfg)
        self.cv_lon_dd_entry.pack(side="left")
        self._bind_coord_pair(self.cv_lat_dd_entry, self.cv_lon_dd_entry,
                              self.cv_lat_dd, self.cv_lon_dd, mode="latlon")
        lf = tk.Frame(geo_dms, bg=PANEL_BG); lf.pack(anchor="w", pady=3)
        tk.Label(lf, text="Latitude",  bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=12, anchor="w").pack(side="left")
        self.cv_lat_dms = DMSEntry(lf, hemispheres=("N","S")); self.cv_lat_dms.pack(side="left")
        lf2 = tk.Frame(geo_dms, bg=PANEL_BG); lf2.pack(anchor="w", pady=3)
        tk.Label(lf2, text="Longitude", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=12, anchor="w").pack(side="left")
        self.cv_lon_dms = DMSEntry(lf2, hemispheres=("E","W")); self.cv_lon_dms.pack(side="left")

        # Cassini N/E input
        self.cv_inp_cas = tk.Frame(inp_host, bg=PANEL_BG)
        self.cv_inp_cas.grid(row=0, column=0, sticky="w")
        self.cv_inp_cas.grid_remove()
        self.cv_mrt48_input_unit = tk.StringVar(value="m")
        self.cv_mrt48_unit_fr = tk.Frame(self.cv_inp_cas, bg=PANEL_BG)
        cv_in_row = tk.Frame(self.cv_mrt48_unit_fr, bg=PANEL_BG)
        cv_in_row.pack(anchor="w")
        tk.Label(cv_in_row, text="Input units :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in _MRT48_UNIT_OPTS:
            tk.Radiobutton(cv_in_row, text=txt, variable=self.cv_mrt48_input_unit,
                           value=val, bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._cv_mrt48_unit_changed).pack(side="left", padx=4)
        self.cv_mrt48_in_hint_lbl = tk.Label(
            cv_in_row, text=mrt48_unit_format_hint("m"), bg=PANEL_BG,
            fg=TEXT_DIM, font=FONT_SM)
        self.cv_mrt48_in_hint_lbl.pack(side="left", padx=(16, 0))
        self.cv_mrt48_unit_fr.pack_forget()
        self.cv_cas_N = tk.StringVar(); self.cv_cas_E = tk.StringVar()
        r_cn = tk.Frame(self.cv_inp_cas, bg=PANEL_BG); r_cn.pack(anchor="w", pady=3)
        self.cv_cas_N_lbl = tk.Label(r_cn, text="Northing  [m]", bg=PANEL_BG, fg=TEXT_MAIN,
                                     font=FONT_LBL, width=22, anchor="w")
        self.cv_cas_N_lbl.pack(side="left")
        self.cv_cas_N_entry = tk.Entry(r_cn, textvariable=self.cv_cas_N, **ecfg)
        self.cv_cas_N_entry.pack(side="left")
        r_ce = tk.Frame(self.cv_inp_cas, bg=PANEL_BG); r_ce.pack(anchor="w", pady=3)
        self.cv_cas_E_lbl = tk.Label(r_ce, text="Easting   [m]", bg=PANEL_BG, fg=TEXT_MAIN,
                                     font=FONT_LBL, width=22, anchor="w")
        self.cv_cas_E_lbl.pack(side="left")
        self.cv_cas_E_entry = tk.Entry(r_ce, textvariable=self.cv_cas_E, **ecfg)
        self.cv_cas_E_entry.pack(side="left")
        self._bind_coord_pair(self.cv_cas_N_entry, self.cv_cas_E_entry,
                              self.cv_cas_N, self.cv_cas_E)

        # UTM input (N/E + zone + hemi)
        self.cv_inp_utm = tk.Frame(inp_host, bg=PANEL_BG)
        self.cv_inp_utm.grid(row=0, column=0, sticky="w")
        self.cv_inp_utm.grid_remove()
        self.cv_utm_N = tk.StringVar(); self.cv_utm_E = tk.StringVar()
        r_un = tk.Frame(self.cv_inp_utm, bg=PANEL_BG); r_un.pack(anchor="w", pady=3)
        tk.Label(r_un, text="Northing  [m]", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        self.cv_utm_N_entry = tk.Entry(r_un, textvariable=self.cv_utm_N, **ecfg)
        self.cv_utm_N_entry.pack(side="left")
        r_ue = tk.Frame(self.cv_inp_utm, bg=PANEL_BG); r_ue.pack(anchor="w", pady=3)
        tk.Label(r_ue, text="Easting   [m]", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        self.cv_utm_E_entry = tk.Entry(r_ue, textvariable=self.cv_utm_E, **ecfg)
        self.cv_utm_E_entry.pack(side="left")
        self._bind_coord_pair(self.cv_utm_N_entry, self.cv_utm_E_entry,
                              self.cv_utm_N, self.cv_utm_E)
        zfr = tk.Frame(self.cv_inp_utm, bg=PANEL_BG); zfr.pack(anchor="w", pady=3)
        tk.Label(zfr, text="UTM Zone :", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL, width=22, anchor="w").pack(side="left")
        self.cv_utm_zone = tk.StringVar()
        tk.Entry(zfr, textvariable=self.cv_utm_zone,
                 **{**ecfg, "width": 6}).pack(side="left", padx=(0,16))
        tk.Label(zfr, text="Hemisphere :", bg=PANEL_BG, fg=TEXT_MAIN,
                 font=FONT_LBL).pack(side="left")
        self.cv_utm_hemi = tk.StringVar(value="N")
        ttk.Combobox(zfr, textvariable=self.cv_utm_hemi, values=["N","S"],
                     width=4, state="readonly", style="H.TCombobox",
                     font=FONT_LBL).pack(side="left")

        # RSO placeholder panel
        self.cv_inp_rso = tk.Frame(inp_host, bg=PANEL_BG)
        self.cv_inp_rso.grid(row=0, column=0, sticky="w")
        self.cv_inp_rso.grid_remove()
        tk.Label(self.cv_inp_rso,
                 text="This panel is unused — RSO MRT48 uses the N/E input above.",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_LBL, justify="left"
                 ).pack(anchor="w", pady=8)

        # ── Buttons ──
        self._cv_kml_lat = None
        self._cv_kml_lon = None
        bf = tk.Frame(p, bg=PANEL_BG); bf.pack(anchor="w", padx=16, pady=8)
        self.cv_btn = self._btn(bf, "▶  Convert", self._do_convert)
        self.cv_btn.pack(side="left", padx=(0,8))
        self._btn(bf, "Clear", self._clr_convert, accent=False).pack(side="left", padx=(0,8))
        self._btn(bf, "🌐 Export KML/KMZ", self._cv_export_kml, accent=False).pack(side="left")

        self._sep(p)

        # ── Output section header + DD/DMS format when To=Geographic ──
        out_hdr = tk.Frame(p, bg=PANEL_BG); out_hdr.pack(fill="x", padx=16, pady=(4, 0))
        tk.Label(out_hdr, text="Output", bg=PANEL_BG, fg=ACCENT2,
                 font=FONT_HDR).pack(side="left")
        self.cv_out_mode_fr = tk.Frame(out_hdr, bg=PANEL_BG)
        self.cv_out_mode = tk.StringVar(value="dd")
        for txt, val in [("Decimal Degrees", "dd"),
                         ("Degrees °   Minutes ′   Seconds ″", "dms")]:
            tk.Radiobutton(self.cv_out_mode_fr, text=txt, variable=self.cv_out_mode, value=val,
                           bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._cv_out_geo_toggle).pack(side="left", padx=4)
        self.cv_out_mode_fr.pack_forget()

        out_host = tk.Frame(p, bg=PANEL_BG); out_host.pack(fill="x", padx=16, pady=4)
        out_host.columnconfigure(0, weight=1)

        self.cv_out_geo = tk.Frame(out_host, bg=PANEL_BG)
        self.cv_out_geo.grid(row=0, column=0, sticky="ew")
        self.cv_out_geo.grid_remove()
        self._cv_out_geo_dd_row = tk.Frame(self.cv_out_geo, bg=PANEL_BG)
        self._cv_out_geo_dd_row.pack(fill="x")
        self.cv_res_lat = ResultCard(self._cv_out_geo_dd_row, "Latitude")
        self.cv_res_lat.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.cv_res_lon = ResultCard(self._cv_out_geo_dd_row, "Longitude")
        self.cv_res_lon.pack(side="left", fill="x", expand=True)
        self._cv_out_geo_dms_row = tk.Frame(self.cv_out_geo, bg=PANEL_BG)
        self.cv_res_lat_dms = ResultCard(self._cv_out_geo_dms_row, "Latitude")
        self.cv_res_lat_dms.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.cv_res_lon_dms = ResultCard(self._cv_out_geo_dms_row, "Longitude")
        self.cv_res_lon_dms.pack(side="left", fill="x", expand=True)
        self._cv_out_geo_dms_row.pack_forget()

        self.cv_out_cas = tk.Frame(out_host, bg=PANEL_BG)
        self.cv_out_cas.grid(row=0, column=0, sticky="ew")
        self.cv_mrt48_output_unit = tk.StringVar(value="m")
        self.cv_mrt48_out_unit_fr = tk.Frame(self.cv_out_cas, bg=PANEL_BG)
        cv_out_u_row = tk.Frame(self.cv_mrt48_out_unit_fr, bg=PANEL_BG)
        cv_out_u_row.pack(anchor="w")
        tk.Label(cv_out_u_row, text="Output units :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in _MRT48_UNIT_OPTS:
            tk.Radiobutton(cv_out_u_row, text=txt, variable=self.cv_mrt48_output_unit,
                           value=val, bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=self._cv_mrt48_output_unit_changed).pack(side="left", padx=4)
        self.cv_mrt48_out_hint_lbl = tk.Label(
            cv_out_u_row, text=mrt48_unit_format_hint("m"), bg=PANEL_BG,
            fg=TEXT_DIM, font=FONT_SM)
        self.cv_mrt48_out_hint_lbl.pack(side="left", padx=(16, 0))
        self.cv_mrt48_out_unit_fr.pack_forget()
        self._cv_out_cas_rf = tk.Frame(self.cv_out_cas, bg=PANEL_BG)
        self._cv_out_cas_rf.pack(fill="x")
        rf = self._cv_out_cas_rf
        self.cv_res_N = ResultCard(rf, "Northing  [m]")
        self.cv_res_N.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.cv_res_E = ResultCard(rf, "Easting   [m]")
        self.cv_res_E.pack(side="left", fill="x", expand=True)

        self.cv_out_utm = tk.Frame(out_host, bg=PANEL_BG)
        self.cv_out_utm.grid(row=0, column=0, sticky="ew")
        self.cv_out_utm.grid_remove()
        rf4 = tk.Frame(self.cv_out_utm, bg=PANEL_BG); rf4.pack(fill="x")
        self.cv_res_utm_N = ResultCard(rf4, "Northing  [m]")
        self.cv_res_utm_N.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.cv_res_utm_E = ResultCard(rf4, "Easting   [m]")
        self.cv_res_utm_E.pack(side="left", fill="x", expand=True)
        rf5 = tk.Frame(self.cv_out_utm, bg=PANEL_BG); rf5.pack(fill="x", pady=(4,0))
        self.cv_res_utm_zone = ResultCard(rf5, "Zone")
        self.cv_res_utm_zone.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.cv_res_utm_hemi = ResultCard(rf5, "Hemisphere")
        self.cv_res_utm_hemi.pack(side="left", fill="x", expand=True)

        self._sep(p)
        self._sec(p, "ECEF / GNSS (Earth-centred XYZ)", fg=ACCENT2)
        ecef_fr = tk.Frame(p, bg=PANEL_BG)
        ecef_fr.pack(fill="x", padx=16, pady=4)
        hf = tk.Frame(ecef_fr, bg=PANEL_BG)
        hf.pack(fill="x", pady=(0, 4))
        tk.Label(hf, text="Ellipsoidal h [m]  (optional — 0 if blank):",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(side="left")
        self.cv_ell_h = tk.StringVar()
        tk.Entry(hf, textvariable=self.cv_ell_h, **self._entry_cfg(width=10)).pack(side="left", padx=8)
        rf_ecef = tk.Frame(ecef_fr, bg=PANEL_BG)
        rf_ecef.pack(fill="x")
        self.cv_res_X = ResultCard(rf_ecef, "X  [m]")
        self.cv_res_X.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.cv_res_Y = ResultCard(rf_ecef, "Y  [m]")
        self.cv_res_Y.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.cv_res_Z = ResultCard(rf_ecef, "Z  [m]")
        self.cv_res_Z.pack(side="left", fill="x", expand=True)
        self.cv_ecef_info = tk.Label(ecef_fr, text="", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.cv_ecef_info.pack(anchor="w", pady=(4, 0))

        self.cv_st = tk.Label(p, text="", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.cv_st.pack(anchor="w", padx=16, pady=6)

        self.cv_from.trace_add("write", lambda *_: self._cv_update())
        self.cv_to.trace_add("write",   lambda *_: self._cv_update())
        self.cv_from_zone.trace_add(
            "write", lambda *_: (self._cv_refresh_datum_info(), self._cv_refresh_output()))
        self.cv_to_zone.trace_add(
            "write", lambda *_: (self._cv_refresh_datum_info(), self._cv_refresh_output()))
        self.cv_from_cas_family.trace_add("write", lambda *_: self._cv_update())
        self.cv_to_cas_family.trace_add("write", lambda *_: self._cv_update())
        self.cv_from_gdm_rev.trace_add("write", lambda *_: self._cv_update())
        self.cv_to_gdm_rev.trace_add("write", lambda *_: self._cv_update())
        self.cv_from_geo_family.trace_add("write", lambda *_: self._cv_update())
        self.cv_to_geo_family.trace_add("write", lambda *_: self._cv_update())
        self.cv_mrt48_input_unit.trace_add("write", lambda *_: self._cv_mrt48_unit_changed())
        self.cv_mrt48_output_unit.trace_add("write", lambda *_: self._cv_mrt48_output_unit_changed())
        self._cv_update()
        self._finish_scrollable_tab(p)

    def _geo_datum(self, side):
        fam = (self.cv_from_geo_family if side == "from"
               else self.cv_to_geo_family).get()
        if fam == "WGS84":
            return "WGS84"
        if fam == "GDM":
            return self._gdm_datum(side)
        if fam == "MRT48":
            return "MRT48"
        return "PMGSN94"

    def _cassini_datum(self, side):
        fam = (self.cv_from_cas_family if side == "from"
               else self.cv_to_cas_family).get()
        if fam == "GDM":
            return (self.cv_from_gdm_rev if side == "from"
                    else self.cv_to_gdm_rev).get()
        return "MRT48"

    def _cv_mrt48_cassini_input_active(self):
        return (self.cv_from.get() == self._CAS
                and self.cv_from_cas_family.get() == "MRT48")

    def _cv_mrt48_input_unit_resolved(self):
        if self._cv_mrt48_cassini_input_active():
            return self.cv_mrt48_input_unit.get()
        return "m"

    def _cv_parse_mrt48_cassini(self, n_str, e_str):
        return mrt48_grid_input_to_m(n_str, e_str, self._cv_mrt48_input_unit_resolved())

    def _cv_mrt48_unit_changed(self, *_):
        if not getattr(self, "cv_cas_N_lbl", None):
            return
        if self._cv_mrt48_cassini_input_active():
            self.cv_mrt48_unit_fr.pack(anchor="w", pady=(0, 4), before=self.cv_cas_N_lbl.master)
            unit = self.cv_mrt48_input_unit.get()
            self.cv_mrt48_in_hint_lbl.config(text=mrt48_unit_format_hint(unit))
            if unit == "links":
                self.cv_cas_N_lbl.config(text="Northing  [links]")
                self.cv_cas_E_lbl.config(text="Easting   [links]")
            else:
                self.cv_cas_N_lbl.config(text="Northing  [m]")
                self.cv_cas_E_lbl.config(text="Easting   [m]")
        else:
            self.cv_mrt48_unit_fr.pack_forget()
            self.cv_cas_N_lbl.config(text="Northing  [m]")
            self.cv_cas_E_lbl.config(text="Easting   [m]")

    def _geoid_mrt48_cassini_input_active(self):
        return (self.g_from.get() == self._CAS
                and self.g_from_cas_family.get() == "MRT48")

    def _geoid_parse_mrt48_cassini(self, n_str, e_str):
        unit = self.g_mrt48_input_unit.get() if self._geoid_mrt48_cassini_input_active() else "m"
        return mrt48_grid_input_to_m(n_str, e_str, unit)

    def _geoid_mrt48_unit_changed(self, *_):
        if not getattr(self, "g_cas_N_lbl", None):
            return
        if self._geoid_mrt48_cassini_input_active():
            self.g_mrt48_unit_fr.pack(anchor="w", pady=(0, 4), before=self.g_cas_N_lbl.master)
            unit = self.g_mrt48_input_unit.get()
            self.g_mrt48_in_hint_lbl.config(text=mrt48_unit_format_hint(unit))
            if unit == "links":
                self.g_cas_N_lbl.config(text="Northing  [links]")
                self.g_cas_E_lbl.config(text="Easting   [links]")
            else:
                self.g_cas_N_lbl.config(text="Northing  [m]")
                self.g_cas_E_lbl.config(text="Easting   [m]")
        else:
            self.g_mrt48_unit_fr.pack_forget()
            self.g_cas_N_lbl.config(text="Northing  [m]")
            self.g_cas_E_lbl.config(text="Easting   [m]")

    def _cv_mrt48_cassini_output_active(self):
        return (self.cv_to.get() == self._CAS
                and self.cv_to_cas_family.get() == "MRT48")

    def _cv_mrt48_output_unit_resolved(self):
        if self._cv_mrt48_cassini_output_active():
            return self.cv_mrt48_output_unit.get()
        return "m"

    def _cv_format_mrt48_cassini(self, n_m, e_m):
        unit = self._cv_mrt48_output_unit_resolved()
        n, e = mrt48_grid_m_to_unit(n_m, e_m, unit)
        return f"{n:,.3f}", f"{e:,.3f}"

    def _cv_set_cassini_out(self, n_m, e_m, dkey):
        if dkey == "MRT48":
            ns, es = self._cv_format_mrt48_cassini(n_m, e_m)
        else:
            ns, es = f"{n_m:,.3f}", f"{e_m:,.3f}"
        self.cv_res_N.set(ns)
        self.cv_res_E.set(es)

    def _cv_mrt48_output_unit_changed(self, *_):
        if not getattr(self, "cv_res_N", None):
            return
        if self._cv_mrt48_cassini_output_active():
            self.cv_mrt48_out_unit_fr.pack(anchor="w", pady=(0, 4), before=self._cv_out_cas_rf)
            unit = self.cv_mrt48_output_unit.get()
            self.cv_mrt48_out_hint_lbl.config(text=mrt48_unit_format_hint(unit))
            if unit == "links":
                self.cv_res_N.set_label("Northing  [links]")
                self.cv_res_E.set_label("Easting   [links]")
            else:
                self.cv_res_N.set_label("Northing  [m]")
                self.cv_res_E.set_label("Easting   [m]")
        else:
            self.cv_mrt48_out_unit_fr.pack_forget()
            self.cv_res_N.set_label("Northing  [m]")
            self.cv_res_E.set_label("Easting   [m]")

    def _gdm_datum(self, side):
        return (self.cv_from_gdm_rev if side == "from"
                else self.cv_to_gdm_rev).get()

    def _cv_datum_blurb(self, sys_name, zone_name=None, side="from"):
        if sys_name == "Geographic (Lat/Lon)":
            d = self._geo_datum(side)
            if d == "MRT48":
                ell = "Modified Everest"
            elif d in ("WGS84", "PMGSN94"):
                ell = "WGS84"
            else:
                ell = "GRS80"
            return f"Datum: {d}   Ellipsoid: {ell}"
        if sys_name == "UTM":
            return "Datum: WGS84   Ellipsoid: WGS84   Standard UTM (k₀=0.9996)"
        if sys_name == "RSO":
            return _rso_datum_blurb(self._rso_zone_resolved(side))
        if sys_name == self._CAS:
            dkey = self._cassini_datum(side)
            if zone_name:
                try:
                    return _zone_datum_blurb(self._find_zone(dkey, zone_name))
                except (StopIteration, ValueError):
                    pass
            ell = "Modified Everest" if dkey == "MRT48" else "GRS80"
            return f"Datum: {dkey}   Ellipsoid: {ell}"
        return ""

    def _cv_refresh_datum_info(self, *_):
        frm = self.cv_from.get()
        to = self.cv_to.get()
        from_zone = self.cv_from_zone.get() if frm == self._CAS else None
        to_zone = self.cv_to_zone.get() if to == self._CAS else None
        self.cv_from_datum_info.config(
            text=self._cv_datum_blurb(frm, from_zone, "from"))
        self.cv_to_datum_info.config(
            text=self._cv_datum_blurb(to, to_zone, "to"))

    def _rso_zone_resolved(self, side="from"):
        if self.cv_rso_family.get() == "MRT48":
            return next(z for z in RSO_ZONES if z["datum"] == "MRT48")
        return next(z for z in RSO_ZONES if z["datum"] == self._gdm_datum(side))

    def _sys_datum(self, sys_name, side="from"):
        if sys_name == "Geographic (Lat/Lon)":
            return self._geo_datum(side)
        if sys_name == self._CAS:
            return self._cassini_datum(side)
        if sys_name == "RSO":
            return ("MRT48" if self.cv_rso_family.get() == "MRT48"
                    else self._gdm_datum(side))
        if sys_name == "UTM":
            return "WGS84"
        return self._cassini_datum(side)

    def _cv_layout_side_options(self, frm, to, frm_cas, to_cas, frm_rso, to_rso,
                                frm_geo, to_geo, rso_gdm):
        """Place Geo / Cassini / RSO datum controls under the From or To column."""
        from_cas_gdm = frm_cas and self.cv_from_cas_family.get() == "GDM"
        to_cas_gdm = to_cas and self.cv_to_cas_family.get() == "GDM"
        from_geo_gdm = frm_geo and self.cv_from_geo_family.get() == "GDM"
        to_geo_gdm = to_geo and self.cv_to_geo_family.get() == "GDM"
        from_needs_gdm_rev = from_cas_gdm or from_geo_gdm or (frm_rso and rso_gdm)
        to_needs_gdm_rev = to_cas_gdm or to_geo_gdm or (to_rso and rso_gdm)

        for fr in (self.cv_from_cas_fr, self.cv_to_cas_fr,
                   self.cv_from_geo_fr, self.cv_to_geo_fr,
                   self.cv_from_gdm_rev_fr, self.cv_to_gdm_rev_fr,
                   self.cv_from_rso_fr, self.cv_to_rso_fr):
            fr.grid_remove()

        from_stack = []
        if frm_geo:
            from_stack.append(self.cv_from_geo_fr)
        if frm_cas:
            from_stack.append(self.cv_from_cas_fr)
        if frm_rso:
            from_stack.append(self.cv_from_rso_fr)
        if from_needs_gdm_rev:
            from_stack.append(self.cv_from_gdm_rev_fr)
        to_stack = []
        if to_geo:
            to_stack.append(self.cv_to_geo_fr)
        if to_cas:
            to_stack.append(self.cv_to_cas_fr)
        if to_rso:
            to_stack.append(self.cv_to_rso_fr)
        if to_needs_gdm_rev:
            to_stack.append(self.cv_to_gdm_rev_fr)

        row = 3
        n = max(len(from_stack), len(to_stack))
        for i in range(n):
            r = row + i
            if i < len(from_stack):
                from_stack[i].grid(row=r, column=0, columnspan=2, sticky="w", pady=(4, 0))
            if i < len(to_stack):
                to_stack[i].grid(row=r, column=3, columnspan=2, sticky="w", pady=(4, 0))
        return row + n

    def _cv_update(self, *_):
        frm = self.cv_from.get()
        to  = self.cv_to.get()
        frm_cas    = frm == self._CAS
        to_cas     = to  == self._CAS
        frm_utm    = frm == "UTM"
        to_utm     = to  == "UTM"
        frm_geo    = frm == "Geographic (Lat/Lon)"
        to_geo     = to  == "Geographic (Lat/Lon)"
        frm_rso    = frm == "RSO"
        to_rso     = to  == "RSO"
        rso_gdm    = self.cv_rso_family.get() == "GDM"

        # FROM sub-selectors
        self.cv_from_state_lbl.grid_remove(); self.cv_from_zone_cb.grid_remove()
        if frm_cas:
            dkey = self._cassini_datum("from")
            zones = self._zones_for(dkey)
            self.cv_from_zone_cb["values"] = zones
            if self.cv_from_zone.get() not in zones:
                self.cv_from_zone.set(zones[0] if zones else "")
            self.cv_from_state_lbl.grid(); self.cv_from_zone_cb.grid()

        # TO sub-selectors
        self.cv_to_state_lbl.grid_remove(); self.cv_to_zone_cb.grid_remove()
        self.cv_to_utm_zone_lbl.grid_remove(); self.cv_to_utm_zone_entry.grid_remove()
        if to_cas:
            dkey = self._cassini_datum("to")
            zones = self._zones_for(dkey)
            self.cv_to_zone_cb["values"] = zones
            if self.cv_to_zone.get() not in zones:
                self.cv_to_zone.set(zones[0] if zones else "")
            self.cv_to_state_lbl.grid(); self.cv_to_zone_cb.grid()
        elif to_utm:
            self.cv_to_utm_zone_lbl.grid(); self.cv_to_utm_zone_entry.grid()

        opt_row = self._cv_layout_side_options(
            frm, to, frm_cas, to_cas, frm_rso, to_rso, frm_geo, to_geo, rso_gdm)
        self._cv_refresh_datum_info()
        self.cv_utm_note.grid_remove()
        if frm_utm or to_utm:
            self.cv_utm_note.grid(row=opt_row, column=0, columnspan=5, sticky="w", pady=(4, 0))

        # Input mode toggle — only visible when From=Geographic
        if frm_geo:
            self.cv_mode_fr.pack(side="right")
        else:
            self.cv_mode_fr.pack_forget()

        # Input panel
        self.cv_inp_geo.grid_remove()
        self.cv_inp_cas.grid_remove()
        self.cv_inp_utm.grid_remove()
        self.cv_inp_rso.grid_remove()
        if frm_geo:        self.cv_inp_geo.grid()
        elif frm_cas or frm_rso: self.cv_inp_cas.grid()
        elif frm_utm:      self.cv_inp_utm.grid()
        self._cv_mrt48_unit_changed()

        # Output panel
        self.cv_out_geo.grid_remove()
        self.cv_out_cas.grid_remove()
        self.cv_out_utm.grid_remove()
        if to_geo:        self.cv_out_geo.grid()
        elif to_cas or to_rso: self.cv_out_cas.grid()
        elif to_utm:      self.cv_out_utm.grid()
        self._cv_mrt48_output_unit_changed()

        if to_geo:
            self.cv_out_mode_fr.pack(side="right", padx=(0, 4))
            self._cv_out_geo_toggle()
        else:
            self.cv_out_mode_fr.pack_forget()

        # Disable Convert if datum transformation required (needs RM160)
        needs_datum_transform = False
        if needs_datum_transform:
            self.cv_btn.config(state="disabled")
            self.cv_st.config(
                text="Coming Soon in Future Update. Stay tuned!  😎",
                fg=ACCENT2)
        else:
            self.cv_btn.config(state="normal")
            self.cv_st.config(text="")

        self._cv_refresh_output()
        self._geoid_refresh_mirror_summary()

    def _cv_clear_output(self):
        for w in (self.cv_res_N, self.cv_res_E, self.cv_res_lat, self.cv_res_lon,
                  self.cv_res_lat_dms, self.cv_res_lon_dms,
                  self.cv_res_utm_N, self.cv_res_utm_E,
                  self.cv_res_utm_zone, self.cv_res_utm_hemi,
                  self.cv_res_X, self.cv_res_Y, self.cv_res_Z):
            w.clear()
        self.cv_ecef_info.config(text="")

    def _cv_input_ready(self):
        frm = self.cv_from.get()
        try:
            if frm == "Geographic (Lat/Lon)":
                if self.cv_mode.get() == "dd":
                    if not self.cv_lat_dd.get().strip() or not self.cv_lon_dd.get().strip():
                        return False
                    float(self.cv_lat_dd.get())
                    float(self.cv_lon_dd.get())
                else:
                    self._cv_get_latlon()
            elif frm in (self._CAS, "RSO"):
                if not self.cv_cas_N.get().strip() or not self.cv_cas_E.get().strip():
                    return False
                float(self.cv_cas_N.get())
                float(self.cv_cas_E.get())
            elif frm == "UTM":
                if not self.cv_utm_N.get().strip() or not self.cv_utm_E.get().strip():
                    return False
                if not self.cv_utm_zone.get().strip():
                    return False
                float(self.cv_utm_N.get())
                float(self.cv_utm_E.get())
                int(self.cv_utm_zone.get())
            else:
                return False
            return True
        except (ValueError, TypeError, tk.TclError):
            return False

    def _cv_refresh_output(self):
        if getattr(self, "_cv_auto_converting", False):
            return
        if self._cv_input_ready():
            self._cv_auto_converting = True
            try:
                self._do_convert()
            finally:
                self._cv_auto_converting = False
        else:
            self._cv_clear_output()

    def _cv_dms_toggle(self):
        if self.cv_mode.get() == "dd":
            self._cv_geo_dms_fr.grid_remove(); self._cv_geo_dd_fr.grid()
        else:
            self._cv_geo_dd_fr.grid_remove(); self._cv_geo_dms_fr.grid()

    def _cv_out_geo_toggle(self):
        if self.cv_out_mode.get() == "dd":
            self._cv_out_geo_dms_row.pack_forget()
            self._cv_out_geo_dd_row.pack(fill="x")
        else:
            self._cv_out_geo_dd_row.pack_forget()
            self._cv_out_geo_dms_row.pack(fill="x", pady=(4, 0))

    def _cv_parse_h(self):
        s = self.cv_ell_h.get().strip()
        return float(s) if s else 0.0

    def _cv_src_datum(self, frm):
        if frm == "UTM":
            return "WGS84"
        if frm == "Geographic (Lat/Lon)":
            return self._geo_datum("from")
        if frm == self._CAS:
            return self._cassini_datum("from")
        if frm == "RSO":
            return self._sys_datum("RSO", "from")
        return self._cassini_datum("from")

    def _cv_dst_datum(self, to, frm=None):
        if to == "UTM":
            return "WGS84"
        if to == "Geographic (Lat/Lon)":
            return self._geo_datum("to")
        if to == self._CAS:
            return self._cassini_datum("to")
        if to == "RSO":
            return self._sys_datum("RSO", "to")
        return self._cassini_datum("to")

    def _cv_output_datum(self, to, geo_datum):
        return geo_datum

    def _cv_shift_geographic(self, lat, lon, src_datum, dst_datum):
        if src_datum == dst_datum or canon_datum(src_datum) == canon_datum(dst_datum):
            return lat, lon
        if not _helmert_key(src_datum, dst_datum):
            raise ValueError(
                f"No JUPEM Helmert path for {src_datum} → {dst_datum} "
                f"(check Malaysia.xml / PKPUP 2021)")
        h = self._cv_parse_h()
        return datum_geo_transform(lat, lon, src_datum, dst_datum, h)

    def _cv_show_ecef(self, lat, lon, datum):
        try:
            h = self._cv_parse_h()
            x, y, z = geo_to_ecef_xyz(lat, lon, datum, h)
            self.cv_res_X.set(f"{x:,.4f}")
            self.cv_res_Y.set(f"{y:,.4f}")
            self.cv_res_Z.set(f"{z:,.4f}")
            ell = ("Mod. Everest" if datum == "MRT48"
                   else "WGS84" if datum in ("WGS84", "PMGSN94") else "GRS80")
            self.cv_ecef_info.config(
                text=f"Datum {datum}  |  {ell} ellipsoid  |  h = {h:.3f} m")
        except Exception:
            for w in (self.cv_res_X, self.cv_res_Y, self.cv_res_Z):
                w.clear()
            self.cv_ecef_info.config(text="")

    def _cv_store_kml_wgs(self, lat, lon, datum):
        """Remember WGS84 position for KML export after a successful convert."""
        try:
            wlat, wlon = self._cv_shift_geographic(lat, lon, datum, "WGS84")
            self._cv_kml_lat, self._cv_kml_lon = wlat, wlon
        except Exception:
            self._cv_kml_lat = self._cv_kml_lon = None

    def _cv_export_kml(self):
        if self._cv_kml_lat is None or self._cv_kml_lon is None:
            messagebox.showwarning(
                "KML Export", "Convert a point first, then export to KML/KMZ.")
            return
        frm, to = self.cv_from.get(), self.cv_to.get()
        name = "Point 1"
        desc = f"{frm} → {to}"
        fp = filedialog.asksaveasfilename(
            title="Export KML / KMZ",
            filetypes=[("KMZ (Google Earth)", "*.kmz"), ("KML", "*.kml")],
            defaultextension=".kmz",
            initialfile="point.kmz")
        if not fp:
            return
        try:
            pts = [(name, self._cv_kml_lat, self._cv_kml_lon, desc)]
            if fp.lower().endswith(".kmz"):
                save_kmz(fp, pts)
            else:
                save_kml(fp, pts)
            self.cv_st.config(
                text=f"✔  Exported to {os.path.basename(fp)}  "
                     f"(WGS84 {self._cv_kml_lat:.6f}°, {self._cv_kml_lon:.6f}°)",
                fg=ACCENT2)
        except Exception as ex:
            messagebox.showerror("Export Error", str(ex))

    def _cv_get_latlon(self):
        if self.cv_mode.get() == "dd":
            return float(self.cv_lat_dd.get()), float(self.cv_lon_dd.get())
        return self.cv_lat_dms.get_dd(), self.cv_lon_dms.get_dd()

    def _input_to_geographic(self, frm):
        """Resolve Convert-tab input to (lat°, lon°, source_datum)."""
        if frm == "Geographic (Lat/Lon)":
            lat, lon = self._cv_get_latlon()
            return lat, lon, self._cv_src_datum(frm)
        if frm == self._CAS:
            dkey = self._cassini_datum("from")
            z = self._find_zone(dkey, self.cv_from_zone.get())
            if dkey == "MRT48":
                n, e = self._cv_parse_mrt48_cassini(self.cv_cas_N.get(), self.cv_cas_E.get())
                lat, lon = mrt48_cassini_to_geo(z, n, e)
            else:
                lat, lon = cassini_to_geo(z, float(self.cv_cas_N.get()), float(self.cv_cas_E.get()))
            return lat, lon, dkey
        if frm == "UTM":
            zn = int(self.cv_utm_zone.get())
            lat, lon = utm_to_geo(WGS84, float(self.cv_utm_N.get()), float(self.cv_utm_E.get()),
                                  zn, self.cv_utm_hemi.get())
            return lat, lon, "WGS84"
        if frm == "RSO":
            rz = self._rso_zone_resolved("from")
            lat, lon = rso_to_geo(rz, float(self.cv_cas_N.get()), float(self.cv_cas_E.get()))
            return lat, lon, self._sys_datum("RSO", "from")
        raise ValueError(f"Unknown input system: {frm}")

    def _do_convert(self):
        frm = self.cv_from.get()
        to  = self.cv_to.get()
        try:
            utm_ok = True

            # Direct MRT48 Cassini ↔ MRT48 RSO (JUPEM polynomial, same state)
            if (frm == self._CAS and self.cv_from_cas_family.get() == "MRT48"
                    and to == "RSO" and self.cv_rso_family.get() == "MRT48"):
                z = self._find_zone("MRT48", self.cv_from_zone.get())
                n, e = self._cv_parse_mrt48_cassini(self.cv_cas_N.get(), self.cv_cas_E.get())
                N, E = cassini_to_mrt48_rso(z, n, e)
                self.cv_res_N.set(f"{N:,.3f}"); self.cv_res_E.set(f"{E:,.3f}")
                lat, lon, _ = self._input_to_geographic(frm)
                self._cv_show_ecef(lat, lon, "MRT48")
                self._cv_store_kml_wgs(lat, lon, "MRT48")
                self.cv_st.config(text=f"✔  {z['name']}" if PUBLIC_BUILD
                                  else f"✔  {z['name']}  |  JUPEM polynomial (m→ch→m)", fg=ACCENT2)
                return
            if (frm == "RSO" and self.cv_rso_family.get() == "MRT48"
                    and to == self._CAS and self.cv_to_cas_family.get() == "MRT48"):
                z = self._find_zone("MRT48", self.cv_to_zone.get())
                N, E = mrt48_rso_to_cassini(z, float(self.cv_cas_N.get()), float(self.cv_cas_E.get()))
                self._cv_set_cassini_out(N, E, "MRT48")
                lat, lon, _ = self._input_to_geographic(frm)
                self._cv_show_ecef(lat, lon, "MRT48")
                self._cv_store_kml_wgs(lat, lon, "MRT48")
                self.cv_st.config(text=f"✔  {z['name']}" if PUBLIC_BUILD
                                  else f"✔  {z['name']}  |  JUPEM polynomial (m→ch→m)", fg=ACCENT2)
                return

            lat, lon, src_datum = self._input_to_geographic(frm)
            if frm == "UTM":
                zn = int(self.cv_utm_zone.get())
                E_in = float(self.cv_utm_E.get())
                utm_ok = self._utm_validate(zn, E_in, lat, lon)

            dst_datum = self._cv_dst_datum(to, frm)
            if dst_datum is None:
                dst_datum = src_datum
            lat, lon = self._cv_shift_geographic(lat, lon, src_datum, dst_datum)

            ecef_datum = self._cv_output_datum(to, dst_datum)
            self._cv_show_ecef(lat, lon, ecef_datum)

            # Produce output from geographic
            if to == "Geographic (Lat/Lon)":
                self._cv_set_geo_out(lat, lon)
                tag = f"  ({dst_datum})" if dst_datum else ""
                via = "" if PUBLIC_BUILD else (
                    "  |  Cassini→RSO→geo" if frm == self._CAS
                    and self.cv_from_cas_family.get() == "MRT48" else "")
                if frm != "UTM" or utm_ok:
                    self.cv_st.config(text=f"✔  Geographic{tag}{via}  (lat={lat:.9f}  lon={lon:.9f})", fg=ACCENT2)
            elif to == self._CAS:
                dkey = self._cassini_datum("to")
                z = self._find_zone(dkey, self.cv_to_zone.get())
                if dkey == "MRT48":
                    N, E = geo_to_mrt48_cassini(z, lat, lon)
                    tag = "" if PUBLIC_BUILD else "via RSO polynomial (geo→RSO→Cassini)"
                else:
                    N, E = geo_to_cassini(z, lat, lon)
                    tag = "" if PUBLIC_BUILD else "Pure Cassini-Soldner"
                self._cv_set_cassini_out(N, E, dkey)
                extra = ""
                if src_datum != dkey and _helmert_key(src_datum, dkey):
                    extra = "" if PUBLIC_BUILD else f"  |  Datum shift {src_datum}→{dkey}"
                if frm != "UTM" or utm_ok:
                    msg = f"✔  {z['name']}" if PUBLIC_BUILD else f"✔  {z['name']}  |  {tag}{extra}"
                    self.cv_st.config(text=msg, fg=ACCENT2)
            elif to == "UTM":
                N, E, zn, hemi = geo_to_utm(WGS84, lat, lon)
                if self.cv_to_utm_override.get().strip():
                    zn = int(self.cv_to_utm_override.get())
                self.cv_res_utm_N.set(f"{N:,.3f}"); self.cv_res_utm_E.set(f"{E:,.3f}")
                self.cv_res_utm_zone.set(str(zn)); self.cv_res_utm_hemi.set(hemi)
                self.cv_st.config(text=f"✔  UTM Zone {zn}{hemi}  |  WGS84", fg=ACCENT2)
            elif to == "RSO":
                rz = self._rso_zone_resolved("to")
                N, E = geo_to_rso(rz, lat, lon)
                self.cv_res_N.set(f"{N:,.3f}"); self.cv_res_E.set(f"{E:,.3f}")
                rso_datum = self._sys_datum("RSO", "to")
                extra = ""
                if src_datum != rso_datum and _helmert_key(src_datum, rso_datum):
                    extra = "" if PUBLIC_BUILD else f"  |  Datum shift {src_datum}→{rso_datum}"
                if frm != "UTM" or utm_ok:
                    msg = (f"✔  {rz['name']}" if PUBLIC_BUILD
                           else f"✔  {rz['name']}  |  Hotine Oblique Mercator{extra}")
                    self.cv_st.config(text=msg, fg=ACCENT2)
            self._cv_store_kml_wgs(lat, lon, ecef_datum)
        except Exception as ex:
            self._cv_kml_lat = self._cv_kml_lon = None
            self.cv_st.config(text=f"⚠  {ex}", fg=ERROR_FG)

    # Peninsular Malaysia bounds (inclusive)
    _MY_LAT  = (1.0,  7.0)
    _MY_LON  = (99.5, 104.7)
    _MY_ZONE = (47, 48)          # only zones 47N and 48N cover the peninsula
    _MY_HEMI = "N"

    def _utm_validate(self, zone, easting, lat, lon):
        warns = []
        # 1. Easting sanity
        if not (100_000 <= easting <= 900_000):
            warns.append(f"Easting {easting:,.0f} m is outside valid UTM range (100 000 – 900 000 m)")
        # 2. Zone must be 47 or 48 for Peninsular Malaysia
        if zone not in self._MY_ZONE:
            warns.append(f"Zone {zone} does not cover Peninsular Malaysia — expected 47 or 48")
        else:
            # 3. Computed longitude must fall inside the declared zone band
            lon_min = (zone - 1) * 6 - 180
            lon_max = zone * 6 - 180
            if not (lon_min <= lon < lon_max):
                suggests = next((z for z in self._MY_ZONE
                                 if (z-1)*6-180 <= lon < z*6-180), None)
                msg = (f"Computed longitude {lon:.4f}°E is outside Zone {zone} "
                       f"({lon_min}°–{lon_max}°E)")
                if suggests:
                    msg += f" — did you mean Zone {suggests}?"
                warns.append(msg)
        # 4. Result must be within Peninsular Malaysia
        if not (self._MY_LAT[0] <= lat <= self._MY_LAT[1]):
            warns.append(f"Computed latitude {lat:.4f}° is outside Peninsular Malaysia "
                         f"({self._MY_LAT[0]}°N – {self._MY_LAT[1]}°N)")
        if not (self._MY_LON[0] <= lon <= self._MY_LON[1]):
            warns.append(f"Computed longitude {lon:.4f}° is outside Peninsular Malaysia "
                         f"({self._MY_LON[0]}°E – {self._MY_LON[1]}°E)")
        if warns:
            self.cv_st.config(text="⚠  " + "  |  ".join(warns), fg=ERROR_FG)
            return False
        return True

    def _cv_set_geo_out(self, lat, lon):
        self.cv_res_lat.set(f"{lat:.9f}")
        self.cv_res_lon.set(f"{lon:.9f}")
        self.cv_res_lat_dms.set(format_dms_full(lat, True))
        self.cv_res_lon_dms.set(format_dms_full(lon, False))

    def _clr_convert(self):
        for v in (self.cv_lat_dd, self.cv_lon_dd, self.cv_cas_N, self.cv_cas_E,
                  self.cv_utm_N, self.cv_utm_E, self.cv_utm_zone):
            v.set("")
        self.cv_lat_dms.clear(); self.cv_lon_dms.clear()
        self.cv_ell_h.set("")
        self._cv_clear_output()
        self._cv_kml_lat = self._cv_kml_lon = None
        self.cv_st.config(text="")

    # ── Batch tab ────────────────────────────────────────────────────
    def _tab_batch_kml(self, p):
        ctrl = tk.Frame(p, bg=PANEL_BG, pady=8)
        ctrl.pack(fill="x", padx=12)
        tk.Label(ctrl, text="From :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).grid(row=0, column=0, sticky="w", padx=(0, 6))
        self.batch_from = tk.StringVar(value=self._CAS)
        ttk.Combobox(ctrl, textvariable=self.batch_from, values=self._CV_SYS,
                     width=22, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).grid(row=0, column=1, sticky="w", padx=(0, 12))
        tk.Label(ctrl, text="→", bg=PANEL_BG, fg=ACCENT,
                 font=("Segoe UI", 12, "bold")).grid(row=0, column=2, padx=8)
        tk.Label(ctrl, text="To :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).grid(row=0, column=3, sticky="w", padx=(0, 6))
        self.batch_to = tk.StringVar(value="Geographic (Lat/Lon)")
        ttk.Combobox(ctrl, textvariable=self.batch_to, values=self._CV_SYS,
                     width=22, state="readonly", style="Zone.TCombobox",
                     font=FONT_LBL).grid(row=0, column=4, sticky="w")

        ctrl2 = tk.Frame(p, bg=PANEL_BG, pady=4)
        ctrl2.pack(fill="x", padx=12)

        self.batch_from_cas_family = tk.StringVar(value="MRT48")
        self.batch_to_cas_family = tk.StringVar(value="MRT48")
        self.batch_from_gdm_rev = tk.StringVar(value="GDM2009")
        self.batch_to_gdm_rev = tk.StringVar(value="GDM2009")
        self.batch_from_geo_family = tk.StringVar(value="WGS84")
        self.batch_to_geo_family = tk.StringVar(value="WGS84")
        self.batch_rso_family = tk.StringVar(value="GDM")
        mrt48_zones = self._zones_for("MRT48")
        z0 = mrt48_zones[0] if mrt48_zones else ""
        self.batch_from_zone = tk.StringVar(value=z0)
        self.batch_to_zone = tk.StringVar(value=z0)

        bchg = self._batch_options_changed
        self.batch_from_cas_fr = self._cv_make_cas_family_fr(
            ctrl2, self.batch_from_cas_family, on_change=bchg)
        self.batch_to_cas_fr = self._cv_make_cas_family_fr(
            ctrl2, self.batch_to_cas_family, on_change=bchg)
        self.batch_from_gdm_rev_fr = self._cv_make_gdm_rev_fr(
            ctrl2, self.batch_from_gdm_rev, on_change=bchg)
        self.batch_to_gdm_rev_fr = self._cv_make_gdm_rev_fr(
            ctrl2, self.batch_to_gdm_rev, on_change=bchg)
        self.batch_from_geo_fr = self._cv_make_geo_datum_fr(
            ctrl2, self.batch_from_geo_family, on_change=bchg)
        self.batch_to_geo_fr = self._cv_make_geo_datum_fr(
            ctrl2, self.batch_to_geo_family, on_change=bchg)
        self.batch_rso_fr = self._cv_make_rso_family_fr(
            ctrl2, self.batch_rso_family, on_change=bchg)

        self.batch_from_zone_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        tk.Label(self.batch_from_zone_fr, text="From state :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        self.batch_from_zone_cb = ttk.Combobox(
            self.batch_from_zone_fr, textvariable=self.batch_from_zone,
            values=mrt48_zones, width=20, state="readonly", style="Zone.TCombobox", font=FONT_SM)
        self.batch_from_zone_cb.pack(side="left")

        self.batch_to_zone_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        tk.Label(self.batch_to_zone_fr, text="To state :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        self.batch_to_zone_cb = ttk.Combobox(
            self.batch_to_zone_fr, textvariable=self.batch_to_zone,
            values=mrt48_zones, width=20, state="readonly", style="Zone.TCombobox", font=FONT_SM)
        self.batch_to_zone_cb.pack(side="left")

        self.batch_mrt48_input_unit = tk.StringVar(value="m")
        self.batch_mrt48_output_unit = tk.StringVar(value="m")
        self.batch_mrt48_unit_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        bin_row = tk.Frame(self.batch_mrt48_unit_fr, bg=PANEL_BG)
        bin_row.pack(anchor="w")
        tk.Label(bin_row, text="MRT48 in :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in _MRT48_UNIT_OPTS:
            tk.Radiobutton(bin_row, text=txt, variable=self.batch_mrt48_input_unit,
                           value=val, bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=bchg).pack(side="left", padx=4)
        self.batch_mrt48_in_hint_lbl = tk.Label(
            bin_row, text=mrt48_unit_format_hint("m"), bg=PANEL_BG,
            fg=TEXT_DIM, font=FONT_SM)
        self.batch_mrt48_in_hint_lbl.pack(side="left", padx=(16, 0))
        self.batch_mrt48_out_unit_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        bout_row = tk.Frame(self.batch_mrt48_out_unit_fr, bg=PANEL_BG)
        bout_row.pack(anchor="w")
        tk.Label(bout_row, text="MRT48 out :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in _MRT48_UNIT_OPTS:
            tk.Radiobutton(bout_row, text=txt, variable=self.batch_mrt48_output_unit,
                           value=val, bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM, command=bchg).pack(side="left", padx=4)
        self.batch_mrt48_out_hint_lbl = tk.Label(
            bout_row, text=mrt48_unit_format_hint("m"), bg=PANEL_BG,
            fg=TEXT_DIM, font=FONT_SM)
        self.batch_mrt48_out_hint_lbl.pack(side="left", padx=(16, 0))

        self.batch_utm_note = tk.Label(ctrl2, text="UTM grid: WGS84 only", bg=PANEL_BG,
                                       fg=TEXT_DIM, font=FONT_SM)

        self.batch_utm_zone = tk.StringVar()
        self.batch_utm_hemi = tk.StringVar(value="N")
        self.batch_utm_out_zone = tk.StringVar()
        self.batch_utm_in_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        utm_in_row = tk.Frame(self.batch_utm_in_fr, bg=PANEL_BG)
        utm_in_row.pack(anchor="w")
        tk.Label(utm_in_row, text="UTM input zone :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        tk.Entry(utm_in_row, textvariable=self.batch_utm_zone,
                 **self._entry_cfg(width=6, font=FONT_SM)).pack(side="left", padx=(0, 12))
        tk.Label(utm_in_row, text="Hemisphere :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        ttk.Combobox(utm_in_row, textvariable=self.batch_utm_hemi, values=["N", "S"],
                     width=4, state="readonly", style="H.TCombobox",
                     font=FONT_SM).pack(side="left", padx=(0, 12))
        tk.Label(utm_in_row, text="(default for rows — or set Zone/Hemi per row in the table)",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(side="left")

        self.batch_utm_out_fr = tk.Frame(ctrl2, bg=PANEL_BG)
        utm_out_row = tk.Frame(self.batch_utm_out_fr, bg=PANEL_BG)
        utm_out_row.pack(anchor="w")
        tk.Label(utm_out_row, text="UTM output zone (auto) :", bg=PANEL_BG, fg=TEXT_DIM,
                 font=FONT_SM).pack(side="left", padx=(0, 6))
        tk.Entry(utm_out_row, textvariable=self.batch_utm_out_zone,
                 **self._entry_cfg(width=6, font=FONT_SM)).pack(side="left", padx=(0, 12))
        tk.Label(utm_out_row, text="leave blank for auto from longitude",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(side="left")

        for var in (self.batch_from, self.batch_to, self.batch_from_cas_family,
                    self.batch_to_cas_family, self.batch_from_gdm_rev, self.batch_to_gdm_rev,
                    self.batch_from_geo_family, self.batch_to_geo_family, self.batch_rso_family,
                    self.batch_mrt48_input_unit, self.batch_mrt48_output_unit):
            var.trace_add("write", self._batch_options_changed)

        self.batch_paste_grid_order = tk.StringVar(value="ne")
        self.batch_paste_order_fr = tk.Frame(p, bg=PANEL_BG)
        self.batch_paste_geo_lbl = tk.Label(
            self.batch_paste_order_fr,
            text="Paste column order: lat/lon auto-detected per row (|value| ≤ 90° = latitude)",
            bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.batch_paste_grid_fr = tk.Frame(self.batch_paste_order_fr, bg=PANEL_BG)
        tk.Label(self.batch_paste_grid_fr, text="Paste column order :", bg=PANEL_BG,
                 fg=TEXT_DIM, font=FONT_SM).pack(side="left", padx=(0, 8))
        for txt, val in (("N then E  (y, x)", "ne"), ("E then N  (x, y)", "en")):
            tk.Radiobutton(self.batch_paste_grid_fr, text=txt,
                           variable=self.batch_paste_grid_order, value=val,
                           bg=PANEL_BG, fg=TEXT_DIM, selectcolor=INPUT_BG,
                           activebackground=PANEL_BG, activeforeground=ACCENT,
                           font=FONT_SM).pack(side="left", padx=4)
        tk.Label(self.batch_paste_grid_fr,
                 text="  — or use a header row (Northing, Easting, …)",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(side="left")

        # ── Toolbar ──
        tb = tk.Frame(p, bg=PANEL_BG, pady=4)
        tb.pack(fill="x", padx=12)
        bst = dict(font=FONT_SM, relief="flat", cursor="hand2", padx=10, pady=4)
        bpri = self._t["btn"]
        tk.Button(tb, text="＋ Add Row",     bg=bpri, fg="white", command=self._batch_add,    **bst).pack(side="left", padx=2)
        tk.Button(tb, text="✕ Delete Row",  bg="#C55A11", fg="white", command=self._batch_del,    **bst).pack(side="left", padx=2)
        tk.Button(tb, text="⟳ Convert All", bg=bpri, fg="white", command=self._batch_convert, **bst).pack(side="left", padx=2)
        tk.Button(tb, text="🗑 Clear All",   bg="#595959", fg="white", command=self._batch_clear,  **bst).pack(side="left", padx=2)
        tk.Button(tb, text="📂 Import TXT/CSV", bg=bpri, fg="white", command=self._batch_import_txt,   **bst).pack(side="left", padx=(16,2))
        tk.Button(tb, text="📂 Import Excel",   bg=bpri, fg="white", command=self._batch_import_xlsx,  **bst).pack(side="left", padx=2)
        tk.Button(tb, text="📋 Paste (Ctrl+V)", bg=bpri, fg="white", command=self._batch_paste,        **bst).pack(side="left", padx=2)
        tk.Button(tb, text="💾 Export Excel",   bg=bpri, fg="white", command=self._batch_export_xlsx,  **bst).pack(side="right", padx=2)
        tk.Button(tb, text="💾 Export CSV",     bg=bpri, fg="white", command=self._batch_export_csv,   **bst).pack(side="right", padx=2)
        tk.Button(tb, text="🌐 Export KML/KMZ", bg="#7B2C2C", fg="white", command=self._batch_export_kml,  **bst).pack(side="right", padx=2)

        # ── Table ──
        tbl_fr = tk.Frame(p, bg=PANEL_BG)
        tbl_fr.pack(fill="both", expand=True, padx=12, pady=(4,4))

        style = ttk.Style()
        self._apply_batch_tree_style()

        self.batch_tree = ttk.Treeview(tbl_fr, style="Batch.Treeview",
                                        show="headings", selectmode="browse")
        vsb = ttk.Scrollbar(tbl_fr, orient="vertical",   command=self.batch_tree.yview)
        hsb = ttk.Scrollbar(tbl_fr, orient="horizontal", command=self.batch_tree.xview)
        self.batch_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.batch_tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tbl_fr.rowconfigure(0, weight=1); tbl_fr.columnconfigure(0, weight=1)

        self.batch_tree.bind("<Double-1>", self._batch_edit)
        self.batch_tree.bind("<Control-v>", lambda e: self._batch_paste())
        self.batch_tree.bind("<Control-V>", lambda e: self._batch_paste())

        # Status
        self.batch_st = tk.Label(p, text="", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.batch_st.pack(anchor="w", padx=12, pady=2)

        self.batch_data = []   # list of dicts
        self._batch_setup_cols()
        self._batch_options_changed()

    def _batch_col_spec(self):
        frm = self.batch_from.get()
        to = self.batch_to.get()
        cols = [("#", "#", 40), ("Name", "name", 100)]
        in_keys = []
        in_labels = ["Name"]
        if frm == "Geographic (Lat/Lon)":
            cols += [("Lat_In", "lat_in", 120), ("Lon_In", "lon_in", 120)]
            in_keys += ["lat_in", "lon_in"]
            in_labels += ["Latitude", "Longitude"]
        elif frm == "UTM":
            cols += [("N_In", "n_in", 110), ("E_In", "e_in", 110),
                     ("Zone_In", "zone_in", 70), ("Hemi_In", "hemi_in", 55)]
            in_keys += ["n_in", "e_in", "zone_in", "hemi_in"]
            in_labels += ["Northing", "Easting", "Zone", "Hemisphere"]
        else:
            unit = ("links" if frm == self._CAS and self.batch_from_cas_family.get() == "MRT48"
                      and self.batch_mrt48_input_unit.get() == "links" else "m")
            suf = "links" if unit == "links" else "m"
            cols += [(f"N_In ({suf})", "n_in", 110), (f"E_In ({suf})", "e_in", 110)]
            in_keys += ["n_in", "e_in"]
            in_labels += [f"Northing ({suf})", f"Easting ({suf})"]
        if to == "Geographic (Lat/Lon)":
            cols += [("Lat_DD", "lat", 120), ("Lon_DD", "lon", 120),
                     ("Lat_DMS", "lat_dms", 200), ("Lon_DMS", "lon_dms", 200)]
        elif to == "UTM":
            cols += [("Northing", "n_out", 110), ("Easting", "e_out", 110),
                     ("Zone", "zone_out", 60), ("Hemi", "hemi_out", 50)]
        else:
            out_suf = ("links" if to == self._CAS and self.batch_to_cas_family.get() == "MRT48"
                         and self.batch_mrt48_output_unit.get() == "links" else "m")
            cols += [(f"Northing ({out_suf})", "n_out", 110), (f"Easting ({out_suf})", "e_out", 110)]
        cols += [("Status", "st", 80)]
        return cols, in_labels, in_keys

    def _batch_setup_cols(self):
        cols_spec, _, _ = self._batch_col_spec()
        cols = [k for _, k, _ in cols_spec]
        self.batch_tree["columns"] = cols
        for hdr, key, w in cols_spec:
            self.batch_tree.heading(key, text=hdr)
            self.batch_tree.column(key, width=w, minwidth=30,
                                   anchor="center" if key not in ("name", "lat_dms", "lon_dms", "st") else "w")
        self._batch_refresh()

    def _batch_geo_datum(self, side):
        fam = (self.batch_from_geo_family if side == "from"
               else self.batch_to_geo_family).get()
        if fam == "WGS84":
            return "WGS84"
        if fam == "GDM":
            return (self.batch_from_gdm_rev if side == "from"
                    else self.batch_to_gdm_rev).get()
        if fam == "MRT48":
            return "MRT48"
        return "PMGSN94"

    def _batch_cassini_datum(self, side):
        fam = (self.batch_from_cas_family if side == "from"
               else self.batch_to_cas_family).get()
        if fam == "GDM":
            return (self.batch_from_gdm_rev if side == "from"
                    else self.batch_to_gdm_rev).get()
        return "MRT48"

    def _batch_parse_mrt48_cassini(self, n, e):
        if (self.batch_from.get() == self._CAS
                and self.batch_from_cas_family.get() == "MRT48"
                and self.batch_mrt48_input_unit.get() == "links"):
            return mrt48_grid_input_to_m(n, e, "links")
        return float(n), float(e)

    def _batch_format_cassini_out(self, n_m, e_m, dkey):
        if dkey == "MRT48" and (self.batch_to.get() == self._CAS
                                and self.batch_to_cas_family.get() == "MRT48"):
            unit = self.batch_mrt48_output_unit.get()
            n, e = mrt48_grid_m_to_unit(n_m, e_m, unit)
            return f"{n:.3f}", f"{e:.3f}"
        return f"{n_m:.3f}", f"{e_m:.3f}"

    def _batch_src_datum(self, frm):
        if frm == "UTM":
            return "WGS84"
        if frm == "Geographic (Lat/Lon)":
            return self._batch_geo_datum("from")
        if frm == self._CAS:
            return self._batch_cassini_datum("from")
        if frm == "RSO":
            return ("MRT48" if self.batch_rso_family.get() == "MRT48"
                    else self.batch_from_gdm_rev.get())
        return self._batch_cassini_datum("from")

    def _batch_dst_datum(self, to, frm=None):
        if to == "UTM":
            return "WGS84"
        if to == "Geographic (Lat/Lon)":
            return self._batch_geo_datum("to")
        if to == self._CAS:
            return self._batch_cassini_datum("to")
        if to == "RSO":
            return ("MRT48" if self.batch_rso_family.get() == "MRT48"
                    else self.batch_to_gdm_rev.get())
        return self._batch_cassini_datum("to")

    def _batch_rso_zone_resolved(self, side):
        if self.batch_rso_family.get() == "MRT48":
            return next(z for z in RSO_ZONES if z["datum"] == "MRT48")
        rev = self.batch_from_gdm_rev if side == "from" else self.batch_to_gdm_rev
        return next(z for z in RSO_ZONES if z["datum"] == rev.get())

    def _batch_refresh_zones(self, side):
        dkey = self._batch_cassini_datum(side)
        zones = self._zones_for(dkey)
        cb = self.batch_from_zone_cb if side == "from" else self.batch_to_zone_cb
        var = self.batch_from_zone if side == "from" else self.batch_to_zone
        cb["values"] = zones
        if var.get() not in zones:
            var.set(zones[0] if zones else "")

    def _batch_options_changed(self, *_):
        frm = self.batch_from.get()
        to = self.batch_to.get()
        for fr in (self.batch_from_cas_fr, self.batch_to_cas_fr,
                   self.batch_from_gdm_rev_fr, self.batch_to_gdm_rev_fr,
                   self.batch_from_geo_fr, self.batch_to_geo_fr,
                   self.batch_rso_fr, self.batch_from_zone_fr, self.batch_to_zone_fr,
                   self.batch_mrt48_unit_fr, self.batch_mrt48_out_unit_fr,
                   self.batch_utm_in_fr, self.batch_utm_out_fr, self.batch_utm_note,
                   self.batch_paste_order_fr):
            fr.pack_forget()

        row = 0
        if frm == self._CAS:
            self.batch_from_cas_fr.pack(fill="x", pady=(2, 0))
            if self.batch_from_cas_family.get() == "GDM":
                self.batch_from_gdm_rev_fr.pack(fill="x", pady=(2, 0))
            else:
                self.batch_mrt48_unit_fr.pack(fill="x", pady=(2, 0))
            self.batch_from_zone_fr.pack(fill="x", pady=(2, 0))
            self._batch_refresh_zones("from")
            row += 1
        elif frm == "Geographic (Lat/Lon)":
            self.batch_from_geo_fr.pack(fill="x", pady=(2, 0))
            if self.batch_from_geo_family.get() == "GDM":
                self.batch_from_gdm_rev_fr.pack(fill="x", pady=(2, 0))
        elif frm == "UTM":
            self.batch_utm_in_fr.pack(fill="x", pady=(2, 0))
            self.batch_utm_note.pack(anchor="w", pady=(2, 0))

        if to == self._CAS:
            self.batch_to_cas_fr.pack(fill="x", pady=(2, 0))
            if self.batch_to_cas_family.get() == "GDM":
                self.batch_to_gdm_rev_fr.pack(fill="x", pady=(2, 0))
            else:
                self.batch_mrt48_out_unit_fr.pack(fill="x", pady=(2, 0))
            self.batch_to_zone_fr.pack(fill="x", pady=(2, 0))
            self._batch_refresh_zones("to")
        elif to == "Geographic (Lat/Lon)":
            self.batch_to_geo_fr.pack(fill="x", pady=(2, 0))
            if self.batch_to_geo_family.get() == "GDM":
                self.batch_to_gdm_rev_fr.pack(fill="x", pady=(2, 0))
        elif to == "UTM":
            self.batch_utm_out_fr.pack(fill="x", pady=(2, 0))
            if frm != "UTM":
                self.batch_utm_note.pack(anchor="w", pady=(2, 0))

        if frm == "RSO" or to == "RSO":
            self.batch_rso_fr.pack(fill="x", pady=(2, 0))
            if self.batch_rso_family.get() == "GDM":
                (self.batch_from_gdm_rev_fr if frm == "RSO" else self.batch_to_gdm_rev_fr
                 ).pack(fill="x", pady=(2, 0))

        self.batch_paste_order_fr.pack(fill="x", padx=12, pady=(4, 0))
        if frm == "Geographic (Lat/Lon)":
            self.batch_paste_grid_fr.pack_forget()
            self.batch_paste_geo_lbl.pack(anchor="w")
        elif frm in (self._CAS, "RSO", "UTM"):
            self.batch_paste_geo_lbl.pack_forget()
            self.batch_paste_grid_fr.pack(anchor="w")
        else:
            self.batch_paste_order_fr.pack_forget()

        if getattr(self, "batch_mrt48_in_hint_lbl", None):
            self.batch_mrt48_in_hint_lbl.config(
                text=mrt48_unit_format_hint(self.batch_mrt48_input_unit.get()))
        if getattr(self, "batch_mrt48_out_hint_lbl", None):
            self.batch_mrt48_out_hint_lbl.config(
                text=mrt48_unit_format_hint(self.batch_mrt48_output_unit.get()))

        self.batch_data.clear()
        self._batch_setup_cols()

    def _batch_row_values(self, i, d):
        cols_spec, _, _ = self._batch_col_spec()
        vals = [i + 1]
        for _, key, _ in cols_spec[1:]:
            vals.append(d.get(key, ""))
        return tuple(vals)

    def _batch_refresh(self):
        for item in self.batch_tree.get_children():
            self.batch_tree.delete(item)
        for i, d in enumerate(self.batch_data):
            tag = "ok" if d.get("st","") == "OK" else ("err" if d.get("st","").startswith("ERR") else "")
            self.batch_tree.insert("", "end", values=self._batch_row_values(i, d), tags=(tag,))
        self.batch_tree.tag_configure("ok",  foreground="#66BB6A")
        self.batch_tree.tag_configure("err", foreground="#EF9A9A")

    def _batch_dialog_widget(self, parent, key, var, ecfg):
        if key in ("zone_in", "zone_out"):
            tk.Entry(parent, textvariable=var,
                     **{**ecfg, "width": 8}).grid(row=0, column=0, sticky="w")
        elif key in ("hemi_in", "hemi_out"):
            ttk.Combobox(parent, textvariable=var, values=["N", "S"],
                         width=4, state="readonly", style="H.TCombobox",
                         font=FONT_LBL).grid(row=0, column=0, sticky="w")
        else:
            tk.Entry(parent, textvariable=var, **ecfg).grid(row=0, column=0, sticky="w")

    def _batch_apply_utm_defaults(self, d):
        if self.batch_from.get() == "UTM":
            if not d.get("zone_in"):
                d["zone_in"] = self.batch_utm_zone.get().strip()
            if not d.get("hemi_in"):
                d["hemi_in"] = self.batch_utm_hemi.get() or "N"

    def _batch_resolved_utm_in(self, d):
        zn_s = (d.get("zone_in") or "").strip() or self.batch_utm_zone.get().strip()
        hemi = (d.get("hemi_in") or "").strip() or self.batch_utm_hemi.get() or "N"
        if not zn_s:
            raise ValueError(
                "UTM zone is required — enter Zone above the table or in each row")
        return int(zn_s), hemi

    def _batch_add(self):
        _, field_labels, field_keys = self._batch_col_spec()
        d = {}
        dlg = tk.Toplevel(self); dlg.title("Add Row"); dlg.configure(bg=DARK_BG)
        dlg.grab_set(); dlg.resizable(False, False)
        ecfg = self._entry_cfg(width=26)
        vars_ = {}
        all_keys = ["name"] + field_keys
        all_labels = ["Name"] + field_labels[1:]
        for i, (lbl, key) in enumerate(zip(all_labels, all_keys)):
            tk.Label(dlg, text=lbl, bg=DARK_BG, fg=TEXT_MAIN,
                     font=FONT_LBL, width=16, anchor="w").grid(row=i, column=0, padx=12, pady=4, sticky="w")
            v = tk.StringVar()
            if key == "zone_in" and self.batch_from.get() == "UTM":
                v.set(self.batch_utm_zone.get())
            elif key == "hemi_in" and self.batch_from.get() == "UTM":
                v.set(self.batch_utm_hemi.get() or "N")
            cell = tk.Frame(dlg, bg=DARK_BG)
            cell.grid(row=i, column=1, padx=(0, 12), pady=4, sticky="w")
            self._batch_dialog_widget(cell, key, v, ecfg)
            vars_[key] = v
            if i == 0:
                dlg.after(50, lambda c=cell: c.winfo_children()[0].focus_set())

        def _save():
            row = {k: v.get().strip() for k, v in vars_.items()}
            self._batch_apply_utm_defaults(row)
            self.batch_data.append(row)
            self._batch_refresh()
            dlg.destroy()
        tk.Button(dlg, text="Add", bg=self._t["btn"], fg="white", font=FONT_LBL,
                  relief="flat", padx=12, pady=5, command=_save
                  ).grid(row=len(all_keys), column=0, columnspan=2, pady=10)
        dlg.bind("<Return>", lambda e: _save())

    def _batch_edit(self, event=None):
        sel = self.batch_tree.selection()
        if not sel: return
        idx = self.batch_tree.index(sel[0])
        d = self.batch_data[idx]
        _, field_labels, field_keys = self._batch_col_spec()
        dlg = tk.Toplevel(self); dlg.title(f"Edit Row {idx+1}"); dlg.configure(bg=DARK_BG)
        dlg.grab_set(); dlg.resizable(False, False)
        ecfg = self._entry_cfg(width=26)
        vars_ = {}
        all_keys = ["name"] + field_keys
        all_labels = ["Name"] + field_labels[1:]
        for i, (lbl, key) in enumerate(zip(all_labels, all_keys)):
            tk.Label(dlg, text=lbl, bg=DARK_BG, fg=TEXT_MAIN,
                     font=FONT_LBL, width=16, anchor="w").grid(row=i, column=0, padx=12, pady=4, sticky="w")
            v = tk.StringVar(value=d.get(key, ""))
            cell = tk.Frame(dlg, bg=DARK_BG)
            cell.grid(row=i, column=1, padx=(0, 12), pady=4, sticky="w")
            self._batch_dialog_widget(cell, key, v, ecfg)
            vars_[key] = v

        def _save():
            for k, v in vars_.items():
                d[k] = v.get().strip()
            self._batch_apply_utm_defaults(d)
            d["st"] = ""
            self._batch_refresh()
            dlg.destroy()
        tk.Button(dlg, text="Save", bg=self._t["btn"], fg="white", font=FONT_LBL,
                  relief="flat", padx=12, pady=5, command=_save
                  ).grid(row=len(all_keys), column=0, columnspan=2, pady=10)
        dlg.bind("<Return>", lambda e: _save())

    def _batch_del(self):
        sel = self.batch_tree.selection()
        if not sel: return
        idx = self.batch_tree.index(sel[0])
        self.batch_data.pop(idx)
        self._batch_refresh()

    def _batch_clear(self):
        if not self.batch_data: return
        if messagebox.askyesno("Clear All", "Clear all batch rows?"):
            self.batch_data.clear(); self._batch_refresh()
            self.batch_st.config(text="")

    def _batch_shift_geographic(self, lat, lon, src, dst):
        if src == dst or canon_datum(src) == canon_datum(dst):
            return lat, lon
        if not _helmert_key(src, dst):
            raise ValueError(
                f"No Helmert path for {src} → {dst} (check Malaysia parameters)")
        return datum_geo_transform(lat, lon, src, dst)

    def _batch_row_to_geographic(self, d, frm):
        if frm == "Geographic (Lat/Lon)":
            return (float(d["lat_in"]), float(d["lon_in"]),
                    self._batch_src_datum(frm))
        if frm == self._CAS:
            dkey = self._batch_cassini_datum("from")
            z = self._find_zone(dkey, self.batch_from_zone.get())
            n, e = self._batch_parse_mrt48_cassini(d["n_in"], d["e_in"])
            if dkey == "MRT48":
                lat, lon = mrt48_cassini_to_geo(z, n, e)
            else:
                lat, lon = cassini_to_geo(z, n, e)
            return lat, lon, dkey
        if frm == "UTM":
            zn, hemi = self._batch_resolved_utm_in(d)
            lat, lon = utm_to_geo(WGS84, float(d["n_in"]), float(d["e_in"]), zn, hemi)
            return lat, lon, "WGS84"
        if frm == "RSO":
            rz = self._batch_rso_zone_resolved("from")
            lat, lon = rso_to_geo(rz, float(d["n_in"]), float(d["e_in"]))
            return lat, lon, self._batch_src_datum(frm)
        raise ValueError(f"Unknown From system: {frm}")

    def _batch_geographic_to_row(self, d, lat, lon, to, dst_datum):
        if to == "Geographic (Lat/Lon)":
            d["lat"] = f"{lat:.9f}"
            d["lon"] = f"{lon:.9f}"
            d["lat_dms"] = format_dms_full(lat, True)
            d["lon_dms"] = format_dms_full(lon, False)
        elif to == self._CAS:
            dkey = self._batch_cassini_datum("to")
            z = self._find_zone(dkey, self.batch_to_zone.get())
            if dkey == "MRT48":
                n, e = geo_to_mrt48_cassini(z, lat, lon)
            else:
                n, e = geo_to_cassini(z, lat, lon)
            ns, es = self._batch_format_cassini_out(n, e, dkey)
            d["n_out"] = ns
            d["e_out"] = es
        elif to == "UTM":
            n, e, zn, hemi = geo_to_utm(WGS84, lat, lon)
            if self.batch_utm_out_zone.get().strip():
                zn = int(self.batch_utm_out_zone.get())
            d["n_out"] = f"{n:.3f}"
            d["e_out"] = f"{e:.3f}"
            d["zone_out"] = str(zn)
            d["hemi_out"] = hemi
        elif to == "RSO":
            rz = self._batch_rso_zone_resolved("to")
            n, e = geo_to_rso(rz, lat, lon)
            d["n_out"] = f"{n:.3f}"
            d["e_out"] = f"{e:.3f}"
        wgs_lat, wgs_lon = self._batch_shift_geographic(lat, lon, dst_datum, "WGS84")
        d["kml_lat"] = f"{wgs_lat:.9f}"
        d["kml_lon"] = f"{wgs_lon:.9f}"

    def _batch_convert_row(self, d):
        frm = self.batch_from.get()
        to = self.batch_to.get()
        if (frm == self._CAS and self.batch_from_cas_family.get() == "MRT48"
                and to == "RSO" and self.batch_rso_family.get() == "MRT48"):
            z = self._find_zone("MRT48", self.batch_from_zone.get())
            n, e = self._batch_parse_mrt48_cassini(d["n_in"], d["e_in"])
            n, e = cassini_to_mrt48_rso(z, n, e)
            d["n_out"] = f"{n:.3f}"
            d["e_out"] = f"{e:.3f}"
            return
        if (frm == "RSO" and self.batch_rso_family.get() == "MRT48"
                and to == self._CAS and self.batch_to_cas_family.get() == "MRT48"):
            z = self._find_zone("MRT48", self.batch_to_zone.get())
            n, e = mrt48_rso_to_cassini(z, float(d["n_in"]), float(d["e_in"]))
            ns, es = self._batch_format_cassini_out(n, e, "MRT48")
            d["n_out"] = ns
            d["e_out"] = es
            return
        lat, lon, src_datum = self._batch_row_to_geographic(d, frm)
        dst_datum = self._batch_dst_datum(to, frm)
        lat, lon = self._batch_shift_geographic(lat, lon, src_datum, dst_datum)
        self._batch_geographic_to_row(d, lat, lon, to, dst_datum)

    def _batch_convert(self):
        ok = err = 0
        for d in self.batch_data:
            try:
                self._batch_convert_row(d)
                d["st"] = "OK"
                ok += 1
            except Exception as ex:
                d["st"] = f"ERR: {ex}"
                err += 1
        self._batch_refresh()
        self.batch_st.config(
            text=f"✔  {ok} converted" + (f"  ⚠  {err} errors" if err else ""),
            fg=ACCENT2 if not err else ERROR_FG)

    def _batch_import_txt(self):
        fp = filedialog.askopenfilename(
            title="Import TXT / CSV",
            filetypes=[("Text/CSV","*.txt *.csv"),("All files","*.*")])
        if not fp: return
        _, field_labels, field_keys = self._batch_col_spec()
        try:
            with open(fp, newline="", encoding="utf-8-sig") as f:
                sample = f.read(2048); f.seek(0)
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t ;")
                reader  = csv.reader(f, dialect)
                rows = list(reader)
            if not rows: return
            # Skip header if first cell looks like text
            start = 1 if rows[0] and not self._is_numeric(rows[0][1] if len(rows[0])>1 else rows[0][0]) else 0
            all_keys = ["name"] + field_keys
            added = 0
            try:
                for row in rows[start:]:
                    cells = [c.strip() for c in row]
                    if not any(cells):
                        continue
                    d = self._batch_row_from_paste_cells(cells, field_keys, all_keys)
                    self.batch_data.append(d)
                    added += 1
            except ValueError as ex:
                messagebox.showerror("Not Lat/Lon", str(ex))
                self.batch_st.config(text="⚠  Import rejected — not lat/lon coordinates.", fg=ERROR_FG)
                return
            self._batch_refresh()
            self.batch_st.config(text=f"✔  Imported {added} rows from {os.path.basename(fp)}", fg=ACCENT2)
        except Exception as ex:
            messagebox.showerror("Import Error", str(ex))

    def _batch_import_xlsx(self):
        fp = filedialog.askopenfilename(
            title="Import Excel",
            filetypes=[("Excel","*.xlsx *.xls"),("All files","*.*")])
        if not fp: return
        _, field_labels, field_keys = self._batch_col_spec()
        try:
            from openpyxl import load_workbook
            wb = load_workbook(fp, read_only=True, data_only=True)
            ws = wb.active
            all_keys = ["name"] + field_keys
            rows = list(ws.iter_rows(values_only=True))
            start = 1 if rows and not self._is_numeric(str(rows[0][1] if len(rows[0])>1 else rows[0][0])) else 0
            added = 0
            null_cells = 0
            try:
                for row in rows[start:]:
                    if not any(c for c in row):
                        continue
                    cells = []
                    for val in row:
                        if val is None:
                            cells.append("")
                            null_cells += 1
                        else:
                            cells.append(str(val).strip())
                    d = self._batch_row_from_paste_cells(cells, field_keys, all_keys)
                    self.batch_data.append(d)
                    added += 1
            except ValueError as ex:
                messagebox.showerror("Not Lat/Lon", str(ex))
                self.batch_st.config(text="⚠  Import rejected — not lat/lon coordinates.", fg=ERROR_FG)
                return
            self._batch_refresh()
            msg = f"✔  Imported {added} rows from {os.path.basename(fp)}"
            if null_cells:
                msg += (f"  ⚠  {null_cells} empty cell(s) detected — "
                        "open and save the Excel file first to flush formula values.")
            self.batch_st.config(text=msg, fg=ACCENT2 if not null_cells else ERROR_FG)
        except Exception as ex:
            messagebox.showerror("Import Error", str(ex))

    def _batch_split_paste_line(self, line):
        if "\t" in line:
            return [c.strip() for c in line.split("\t")]
        for sep in (",", ";"):
            if sep in line:
                parts = [p.strip() for p in line.split(sep)]
                if len(parts) >= 2:
                    return parts
        return [p.strip() for p in line.split() if p.strip()]

    def _batch_apply_paste_coord_order(self, d, field_keys):
        if self.batch_from.get() != "Geographic (Lat/Lon)":
            return
        if d.get("lat_in") and d.get("lon_in"):
            err = validate_lat_lon_paste((d["lat_in"], d["lon_in"]))
            if err:
                raise ValueError(err)
            lat, lon = assign_lat_lon((d["lat_in"], d["lon_in"]))
            d["lat_in"], d["lon_in"] = lat, lon

    def _batch_paste_has_leading_name(self, cells, field_keys, all_keys):
        """True when the first pasted column is a point name, not a coordinate."""
        if not cells:
            return False
        if len(cells) >= len(all_keys):
            return True
        first = cells[0].strip().replace(",", "")
        if self._is_numeric(first):
            return False
        return len(cells) >= 2

    def _batch_assign_paste_data_cells(self, d, data_cells, field_keys):
        order_keys = list(field_keys)
        if ("n_in" in order_keys and "e_in" in order_keys
                and self.batch_paste_grid_order.get() == "en"):
            i_n, i_e = order_keys.index("n_in"), order_keys.index("e_in")
            order_keys[i_n], order_keys[i_e] = "e_in", "n_in"
        for i, key in enumerate(order_keys):
            if i < len(data_cells):
                d[key] = data_cells[i].strip()

    def _batch_row_from_paste_cells(self, cells, field_keys, all_keys, header_map=None):
        d = {k: "" for k in all_keys}
        if header_map is not None:
            for key in all_keys:
                idx = header_map.get(key)
                if idx is not None and idx < len(cells):
                    d[key] = cells[idx].strip()
            self._batch_apply_paste_coord_order(d, field_keys)
            self._batch_apply_utm_defaults(d)
            return d

        if self._batch_paste_has_leading_name(cells, field_keys, all_keys):
            d["name"] = cells[0].strip()
            data_cells = cells[1:]
        else:
            data_cells = cells

        self._batch_assign_paste_data_cells(d, data_cells, field_keys)
        self._batch_apply_paste_coord_order(d, field_keys)
        self._batch_apply_utm_defaults(d)
        return d

    def _batch_paste(self):
        try:
            text = self.clipboard_get()
        except tk.TclError:
            self.batch_st.config(text="⚠  Clipboard is empty.", fg=ERROR_FG)
            return
        _, field_labels, field_keys = self._batch_col_spec()
        all_keys = ["name"] + field_keys
        lines = [l for l in text.splitlines() if l.strip()]
        if not lines:
            self.batch_st.config(text="⚠  Nothing to paste.", fg=ERROR_FG)
            return
        parsed = [self._batch_split_paste_line(ln) for ln in lines]
        start = 0
        header_map = None
        if parsed and _batch_is_header_row(parsed[0]):
            header_map = _batch_header_col_map(parsed[0], all_keys)
            if _batch_header_usable(header_map, field_keys):
                start = 1
            else:
                header_map = None
        added = 0
        new_rows = []
        try:
            for cells in parsed[start:]:
                if not any(c.strip() for c in cells):
                    continue
                new_rows.append(
                    self._batch_row_from_paste_cells(cells, field_keys, all_keys, header_map))
        except ValueError as ex:
            messagebox.showerror("Not Lat/Lon", str(ex))
            self.batch_st.config(text="⚠  Paste rejected — not lat/lon coordinates.", fg=ERROR_FG)
            return
        for d in new_rows:
            self.batch_data.append(d)
            added += 1
        self._batch_refresh()
        self.batch_st.config(text=f"✔  Pasted {added} rows from clipboard.", fg=ACCENT2)

    def _is_numeric(self, s):
        try: float(s); return True
        except: return False

    def _batch_export_xlsx(self):
        if not self.batch_data:
            messagebox.showinfo("Export", "No data to export."); return
        fp = filedialog.asksaveasfilename(defaultextension=".xlsx",
            filetypes=[("Excel","*.xlsx")], initialfile="batch_export.xlsx")
        if not fp: return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook(); ws = wb.active; ws.title = "Batch Convert"
            cols_spec, _, _ = self._batch_col_spec()
            headers = [h for h,_,_ in cols_spec]
            for ci, h in enumerate(headers, 1):
                c = ws.cell(row=1, column=ci, value=h)
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", start_color="1F4E79")
                c.alignment = Alignment(horizontal="center")
            for ri, d in enumerate(self.batch_data, 2):
                for ci, val in enumerate(self._batch_row_values(ri-2, d), 1):
                    ws.cell(row=ri, column=ci, value=val)
            wb.save(fp)
            os.startfile(fp)
        except Exception as ex:
            messagebox.showerror("Export Error", str(ex))

    def _batch_export_csv(self):
        if not self.batch_data:
            messagebox.showinfo("Export", "No data to export."); return
        fp = filedialog.asksaveasfilename(defaultextension=".csv",
            filetypes=[("CSV","*.csv")], initialfile="batch_export.csv")
        if not fp: return
        try:
            cols_spec, _, _ = self._batch_col_spec()
            headers = [h for h,_,_ in cols_spec]
            with open(fp, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f)
                w.writerow(headers)
                for i, d in enumerate(self.batch_data):
                    w.writerow(self._batch_row_values(i, d))
            self.batch_st.config(text=f"✔  Exported to {os.path.basename(fp)}", fg=ACCENT2)
        except Exception as ex:
            messagebox.showerror("Export Error", str(ex))

    def _batch_export_kml(self):
        if not self.batch_data:
            messagebox.showinfo("Export", "No data to export."); return
        points = []
        for i, d in enumerate(self.batch_data):
            try:
                if d.get("kml_lat") and d.get("kml_lon"):
                    lat = float(d["kml_lat"])
                    lon = float(d["kml_lon"])
                elif d.get("lat") and d.get("lon"):
                    lat = float(d["lat"])
                    lon = float(d["lon"])
                elif d.get("lat_in") and d.get("lon_in"):
                    lat = float(d["lat_in"])
                    lon = float(d["lon_in"])
                else:
                    continue
                name = d.get("name", "") or f"Point {i+1}"
                points.append((name, lat, lon, f"Row {i+1}"))
            except ValueError:
                pass
        if not points:
            messagebox.showwarning("KML Export", "No valid geographic coordinates to export.\nRun Convert All first."); return
        fp = filedialog.asksaveasfilename(
            title="Export KML / KMZ",
            filetypes=[("KMZ (Google Earth)","*.kmz"),("KML","*.kml")],
            defaultextension=".kmz", initialfile="batch_export.kmz")
        if not fp: return
        try:
            if fp.lower().endswith(".kmz"):
                save_kmz(fp, points)
            else:
                save_kml(fp, points)
            self.batch_st.config(text=f"✔  Exported {len(points)} points to {os.path.basename(fp)}", fg=ACCENT2)
        except Exception as ex:
            messagebox.showerror("Export Error", str(ex))

    # ── Tab 3: Geoid Height ────────────────────────────────────────────────
    def _tab_geoid(self, p):
        self._geoid_scroll_inner = p
        self._geoid_build_input_section(p)
        self._sep(p)
        self._sec(p, "WGeoid04  —  Malaysia Geoid Undulation (N)")

        # Geoid grid (embedded in .exe — hide file path from students)
        hint_fr = tk.Frame(p, bg=PANEL_BG)
        hint_fr.pack(anchor="w", padx=16, pady=(0,4))
        from app_paths import GEOID_GSF_NAME, geoid_file_path
        self.gff_var = tk.StringVar(value=geoid_file_path())
        if is_student():
            tk.Label(hint_fr,
                     text="Geoid model is built into this application.",
                     bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(anchor="w")
        else:
            tk.Label(hint_fr, text=f"Geoid grid ({GEOID_GSF_NAME}):",
                     bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM).pack(side="left")
            ecfg = self._entry_cfg(width=54, font=FONT_SM)
            tk.Entry(hint_fr, textvariable=self.gff_var, **ecfg).pack(side="left", padx=8)

        # ── Height inputs ──
        self._sep(p)
        self._sec(p, "Height Input")
        ht_fr = tk.Frame(p, bg=PANEL_BG)
        ht_fr.pack(anchor="w", padx=16, pady=4)
        ecfg2 = self._entry_cfg(width=22)
        self.g_h  = tk.StringVar()   # ellipsoidal
        self.g_H  = tk.StringVar()   # orthometric
        for var, lbl, hint in [
                (self.g_h,  "Ellipsoidal height  h  [m]",  "leave blank to get N only"),
                (self.g_H,  "Orthometric height  H  [m]",  "leave blank to get N only")]:
            r = tk.Frame(ht_fr, bg=PANEL_BG)
            r.pack(anchor="w", pady=2)
            tk.Label(r, text=lbl, bg=PANEL_BG, fg=TEXT_MAIN,
                     font=FONT_LBL, width=26, anchor="w").pack(side="left")
            tk.Entry(r, textvariable=var, **ecfg2).pack(side="left")
            tk.Label(r, text=hint, bg=PANEL_BG, fg=TEXT_DIM,
                     font=FONT_SM).pack(side="left", padx=8)
        tk.Label(ht_fr,
                 text="Tip: when mirroring Convert, ellipsoidal h can also be taken from Convert → h if left blank.",
                 bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM, wraplength=560,
                 justify="left").pack(anchor="w", pady=(6, 0))

        # ── Buttons ──
        bf = tk.Frame(p, bg=PANEL_BG); bf.pack(anchor="w", padx=16, pady=10)
        self._btn(bf, "▶  Compute", self._do_geoid).pack(side="left", padx=(0,8))
        self._btn(bf, "Clear", self._clr_geoid, accent=False).pack(side="left")

        self._sep(p)
        self._sec(p, "Results", fg=ACCENT2)
        rf = tk.Frame(p, bg=PANEL_BG); rf.pack(fill="x", padx=16, pady=4)
        self.g_N_out  = ResultCard(rf, "Geoid undulation  N  [m]")
        self.g_N_out.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.g_H_out  = ResultCard(rf, "Orthometric  H = h − N  [m]")
        self.g_H_out.pack(side="left", fill="x", expand=True)

        rf2 = tk.Frame(p, bg=PANEL_BG); rf2.pack(fill="x", padx=16, pady=4)
        self.g_h_out  = ResultCard(rf2, "Ellipsoidal  h = H + N  [m]")
        self.g_h_out.pack(side="left", fill="x", expand=True, padx=(0,8))
        self.g_latlon_out = ResultCard(rf2, "Used Lat / Lon")
        self.g_latlon_out.pack(side="left", fill="x", expand=True)

        rf3 = tk.Frame(p, bg=PANEL_BG); rf3.pack(fill="x", padx=16, pady=4)
        self.g_X_out = ResultCard(rf3, "ECEF  X  [m]")
        self.g_X_out.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.g_Y_out = ResultCard(rf3, "ECEF  Y  [m]")
        self.g_Y_out.pack(side="left", fill="x", expand=True, padx=(0, 8))
        self.g_Z_out = ResultCard(rf3, "ECEF  Z  [m]")
        self.g_Z_out.pack(side="left", fill="x", expand=True)

        self.g_st = tk.Label(p, text="", bg=PANEL_BG, fg=TEXT_DIM, font=FONT_SM)
        self.g_st.pack(anchor="w", padx=16, pady=4)
        self._finish_scrollable_tab(p)

    def _do_geoid(self):
        try:
            gff = self.gff_var.get().strip()
            if self.geoid_mirror_convert.get():
                lat, lon, datum = self._input_to_geographic(self.cv_from.get())
            else:
                lat, lon, datum = self._geoid_input_to_geographic()

            N, err = geoid_undulation(lat, lon, gff)
            if N is None:
                self.g_st.config(text=f"⚠  {err}", fg=ERROR_FG)
                for w in (self.g_N_out, self.g_H_out, self.g_h_out, self.g_latlon_out,
                          self.g_X_out, self.g_Y_out, self.g_Z_out):
                    w.clear()
                return

            self.g_N_out.set(f"{N:.4f}")
            self.g_latlon_out.set(f"{lat:.6f}°N   {lon:.6f}°E")

            h_str = self.g_h.get().strip()
            if not h_str and self.geoid_mirror_convert.get():
                h_str = self.cv_ell_h.get().strip()
            H_str = self.g_H.get().strip()
            h_ecef = float(h_str) if h_str else (float(H_str) + N if H_str else 0.0)
            x, y, zc = geo_to_ecef_xyz(lat, lon, datum, h_ecef)
            self.g_X_out.set(f"{x:,.4f}")
            self.g_Y_out.set(f"{y:,.4f}")
            self.g_Z_out.set(f"{zc:,.4f}")

            if h_str:
                h = float(h_str)
                self.g_H_out.set(f"{h - N:.4f}")
            else:
                self.g_H_out.set("—  (enter h above)")

            if H_str:
                H = float(H_str)
                self.g_h_out.set(f"{H + N:.4f}")
            else:
                self.g_h_out.set("—  (enter H above)")

            self.g_st.config(
                text=f"✔  WGeoid04 bilinear interpolation  |  N = {N:.4f} m  |  Coverage: Lat 0–8°N, Lon 98–107°E",
                fg=ACCENT2)
        except ValueError:
            self.g_st.config(text="⚠  Enter valid numeric values.", fg=ERROR_FG)
        except Exception as ex:
            self.g_st.config(text=f"⚠  {ex}", fg=ERROR_FG)

    def _clr_geoid(self):
        if not self.geoid_mirror_convert.get():
            for v in (self.g_lat_dd, self.g_lon_dd, self.g_cas_N, self.g_cas_E,
                      self.g_utm_N, self.g_utm_E, self.g_utm_zone):
                v.set("")
            self.g_lat_dms.clear()
            self.g_lon_dms.clear()
        for v in (self.g_h, self.g_H):
            v.set("")
        for w in (self.g_N_out, self.g_H_out, self.g_h_out, self.g_latlon_out,
                  self.g_X_out, self.g_Y_out, self.g_Z_out):
            w.clear()
        self.g_st.config(text="")

    # ── Tab: Distribute (lecturer edition only) ─────────────────────────────
    def _tab_distribute(self, p):
        from app_paths import GEOID_GSF_NAME, geoid_file_path, resource as _res, exe_dir, bundled_path

        card = tk.Frame(p, bg=CARD_BG, padx=28, pady=24)
        card.pack(fill="both", expand=True, padx=16, pady=16)

        tk.Label(card, text="Distribute to Students",
                 bg=CARD_BG, fg="#C9A227", font=("Segoe UI", 14, "bold")
                 ).pack(anchor="w", pady=(0, 8))
        tk.Label(card,
                 text="Export a ready-to-share student package (.zip).\n"
                      "Contains the student .exe and README only — geoid grid is already inside the .exe.\n"
                      "Students double-click the .exe; no separate data files.",
                 bg=CARD_BG, fg=TEXT_MAIN, font=FONT_SM, justify="left",
                 wraplength=620).pack(anchor="w", pady=(0, 16))

        xml_p = _res("Malaysia.xml")
        st_txt = f"Malaysia.xml : {xml_p}" if os.path.isfile(xml_p) else "Malaysia.xml : not found (place beside lecturer exe)"
        self.dist_st = tk.Label(card, text=st_txt, bg=CARD_BG, fg=TEXT_DIM, font=FONT_SM)
        self.dist_st.pack(anchor="w", pady=(0, 12))

        from student_pack import describe_student_template
        self.dist_template_lbl = tk.Label(
            card, text=describe_student_template(), bg=CARD_BG, fg=TEXT_DIM,
            font=FONT_SM, justify="left", wraplength=620)
        self.dist_template_lbl.pack(anchor="w", pady=(0, 12))

        nf = tk.Frame(card, bg=CARD_BG)
        nf.pack(anchor="w", pady=(0, 12))
        tk.Label(nf, text="Your name (optional — From your lecturer name in student README):",
                 bg=CARD_BG, fg=TEXT_MAIN, font=FONT_SM).pack(anchor="w")
        self.dist_lecturer_name = tk.StringVar()
        tk.Entry(nf, textvariable=self.dist_lecturer_name,
                 **self._entry_cfg(width=48, font=FONT_LBL)).pack(anchor="w", pady=(4, 0))

        bf = tk.Frame(card, bg=CARD_BG)
        bf.pack(anchor="w", pady=4)
        self._btn(bf, "Export Student ZIP…", self._export_student_zip).pack(side="left", padx=(0, 8))
        self._btn(bf, "Generate class test coordinates…", self._export_class_tests, accent=False).pack(side="left")

        tk.Label(card,
                 text="Tip: custom logo/icon → bin/branding/lecturer/  "
                      "(developer rebuild only; Distribute works from this .exe).",
                 bg=CARD_BG, fg=TEXT_DIM, font=FONT_SM, justify="left"
                 ).pack(anchor="w", pady=(20, 0))

    def _export_student_zip(self):
        from student_pack import (
            create_student_zip,
            describe_student_template,
            resolve_student_exe_path,
        )

        dest = filedialog.asksaveasfilename(
            title="Save student package",
            defaultextension=".zip",
            initialfile="UTM_Coordinate_Wizard_Students.zip",
            filetypes=[("ZIP archive", "*.zip")])
        if not dest:
            return
        try:
            self.dist_st.config(text="Preparing student package…", fg=TEXT_DIM)
            self.update_idletasks()
            template = resolve_student_exe_path(rebuild_if_stale=True)
            if getattr(self, "dist_template_lbl", None):
                self.dist_template_lbl.config(text=describe_student_template())
            create_student_zip(
                dest, template, self.dist_lecturer_name.get().strip())
            from student_pack import UTM_STUDENT_EXE_NAME
            self.dist_st.config(
                text=f"✔  Student package saved: {dest}\n    from: {template}",
                fg=ACCENT2)
            messagebox.showinfo(
                "Student package",
                f"Saved:\n{dest}\n\n"
                f"Student exe: {UTM_STUDENT_EXE_NAME}\n"
                "Share this zip with your class.")
        except Exception as ex:
            self.dist_st.config(text=f"⚠  Export failed", fg=ERROR_FG)
            messagebox.showerror("Export failed", str(ex))

    def _export_class_tests(self):
        import random
        dest = filedialog.asksaveasfilename(
            title="Save class test coordinates",
            defaultextension=".txt",
            initialfile="Class Test Coordinates MRT48 Cassini to RSO.txt",
            filetypes=[("Text", "*.txt")])
        if not dest:
            return
        random.seed(77)
        lines = [
            "=" * 72,
            "  CLASS TEST — MRT48 Cassini → MRT48 RSO",
            f"  Generated by {UTM_APP_NAME} (Lecturer Edition)",
            "=" * 72,
            "",
        ]
        states = [
            ("Johor", "Johor", 1.2, 2.75, 102.3, 104.0),
            ("Perak Utara", "Perak Utara", 3.7, 5.7, 100.3, 101.9),
            ("Kedah", "Kedah", 5.6, 6.7, 99.6, 100.75),
        ]
        for label, kw, la0, la1, lo0, lo1 in states:
            zone = next((z for z in ZONES if z["datum"] == "MRT48" and kw in z["name"]), None)
            if not zone:
                continue
            lines += ["-" * 72, f"  {label}", "-" * 72,
                      f"  {'No':<4} {'Cassini N':>14} {'Cassini E':>14}   {'RSO N':>12} {'RSO E':>12}"]
            for i in range(4):
                lat = round(random.uniform(la0, la1), 6)
                lon = round(random.uniform(lo0, lo1), 6)
                cN, cE = geo_to_cassini(zone, lat, lon)
                rN, rE = cassini_to_mrt48_rso(zone, cN, cE)
                lines.append(f"  {i+1:<4} {cN:>14.4f} {cE:>14.4f}   {rN:>12.4f} {rE:>12.4f}")
            lines.append("")
        lines.append("=" * 72)
        with open(dest, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        self.dist_st.config(text=f"✔  Class tests saved: {dest}", fg=ACCENT2)
        messagebox.showinfo("Class tests", f"Saved:\n{dest}")

    # ── Tab 4: MRT48 ↔ GDM ────────────────────────────────────────────────
    def _tab_datum(self, p):
        outer = tk.Frame(p, bg=PANEL_BG)
        outer.pack(fill="both", expand=True)
        outer.grid_rowconfigure(0, weight=1)
        outer.grid_columnconfigure(0, weight=1)

        card = tk.Frame(outer, bg=CARD_BG, padx=32, pady=32)
        card.grid(row=0, column=0)

        tk.Label(card, text="MRT48  ↔  GDM2000 / GDM2009",
                 bg=CARD_BG, fg=ACCENT, font=("Segoe UI", 14, "bold")
                 ).pack(pady=(0, 16))

        tk.Label(card,
                 text="Datum transformation parameters are required for this conversion.\n"
                      "These parameters are published by and available from:\n",
                 bg=CARD_BG, fg=TEXT_MAIN, font=FONT_LBL, justify="center"
                 ).pack()

        tk.Label(card,
                 text="Jabatan Ukur dan Pemetaan Malaysia (JUPEM)",
                 bg=CARD_BG, fg=ACCENT2, font=("Segoe UI", 11, "bold")
                 ).pack()

        tk.Label(card,
                 text="Department of Survey and Mapping Malaysia",
                 bg=CARD_BG, fg=TEXT_DIM, font=FONT_SM
                 ).pack(pady=(0, 20))

        tk.Label(card,
                 text="Once the official JUPEM transformation parameters are obtained,\n"
                      "this tab will be activated for MRT48 ↔ GDM coordinate conversion.",
                 bg=CARD_BG, fg=TEXT_DIM, font=FONT_SM, justify="center"
                 ).pack()

    # ── Tab 5 ──────────────────────────────────────────────────────────────
    def _tab_help(self, p):
        fr = tk.Frame(p, bg=PANEL_BG); fr.pack(fill="both", expand=True, padx=4, pady=4)
        sb = tk.Scrollbar(fr); sb.pack(side="right", fill="y")
        txt = tk.Text(fr, bg=PANEL_BG, fg=TEXT_MAIN, font=("Consolas",9),
                      relief="flat", padx=16, pady=12, wrap="word",
                      yscrollcommand=sb.set, width=74, height=22)
        txt.pack(fill="both", expand=True)
        sb.config(command=txt.yview)
        txt.insert("1.0", _help_text())
        txt.config(state="disabled")


# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    # UTM lecturer/student — set EA_EDITION or use 5_UTM_Lecturer_Dev.bat / frozen .exe.
    # EA original edition: bin/ea_coordinate_wizard.py (5_EA_Coordinate_Wizard.bat).
    CassiniApp().mainloop()
